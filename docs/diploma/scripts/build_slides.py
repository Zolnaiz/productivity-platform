# -*- coding: utf-8 -*-
"""I үзлэгийн илтгэл (.pptx): python scripts/build_slides.py

Canva дээрх илтгэлтэй ижил өнгө, бүтэцтэй. Агуулга нь latex/ доторх дипломын
бичвэрээс (шаардлага, use-case, мөрдөх матриц, KPI, алгоритм, ном зүй) авсан тул
тоо, нэр томьёог тэндээс өөрчилвөл энд ч засна.
"""
import os

from lxml import etree
from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Cm, Emu, Pt

HERE = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
FIG = os.path.join(HERE, 'fig')

NAVY = RGBColor(0x16, 0x20, 0x2C)
NAVY2 = RGBColor(0x1E, 0x2B, 0x39)
YELLOW = RGBColor(0xF5, 0xC5, 0x18)
INK = RGBColor(0x20, 0x2A, 0x35)
MUTED = RGBColor(0x65, 0x71, 0x7D)
CARD = RGBColor(0xF3, 0xF5, 0xF7)
LINE = RGBColor(0xD8, 0xDE, 0xE4)
GREEN = RGBColor(0x2F, 0x7D, 0x68)
AMBER = RGBColor(0xA8, 0x6B, 0x0C)
RED = RGBColor(0xB8, 0x50, 0x42)
BLUE = RGBColor(0x2F, 0x5D, 0xA8)
GREY = RGBColor(0x9A, 0xA5, 0xB1)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
PALE = RGBColor(0xCB, 0xD3, 0xDA)
TINT = RGBColor(0xFE, 0xF6, 0xD8)
CODE_INK = RGBColor(0xE6, 0xED, 0xF3)
CODE_NOTE = RGBColor(0x8B, 0x98, 0xA5)
FONT = 'Arial'
MONO = 'Consolas'
SYMBOL = 'Segoe UI Symbol'

AUTHOR = 'Б.Билгүүн'
TITLE = 'Байгууллагын бүтээмжийг дэмжих платформ хөгжүүлэлт'

prs = Presentation()
prs.slide_width, prs.slide_height = Cm(33.867), Cm(19.05)
W, H = prs.slide_width, prs.slide_height
M = Cm(1.6)
CW = W - 2 * M
blank = prs.slide_layouts[6]
prs.core_properties.title = TITLE
prs.core_properties.author = AUTHOR


# --------------------------------------------------------------------------- helpers

def text(s, x, y, w, h, lines, size=16, color=INK, bold=False, align=PP_ALIGN.LEFT, font=FONT,
         spacing=None, after=4, line=1.08, anchor=MSO_ANCHOR.TOP, italic=False):
    """Бичвэр. lines: мөрүүд; мөр бүр str, (str, style) эсвэл тэдгээрийн жагсаалт (run-ууд)."""
    box = s.shapes.add_textbox(int(x), int(y), int(w), int(h))
    frame = box.text_frame
    frame.word_wrap = True
    frame.margin_left = frame.margin_right = frame.margin_top = frame.margin_bottom = 0
    frame.vertical_anchor = anchor
    if not isinstance(lines, list):
        lines = [lines]
    for i, item in enumerate(lines):
        p = frame.paragraphs[0] if i == 0 else frame.add_paragraph()
        p.alignment = align
        p.space_after = Pt(after)
        p.line_spacing = line
        for part in (item if isinstance(item, list) else [item]):
            value, style = part if isinstance(part, tuple) else (part, {})
            run = p.add_run()
            run.text = value
            f = run.font
            f.name = style.get('font', font)
            f.size = Pt(style.get('size', size))
            f.bold = style.get('bold', bold)
            f.italic = style.get('italic', italic)
            f.color.rgb = style.get('color', color)
            gap = style.get('spacing', spacing)
            if gap:
                run._r.get_or_add_rPr().set('spc', str(int(gap * 100)))
    return box


def rect(s, x, y, w, h, fill, radius=None, line=None, line_w=Pt(1), shape=None):
    kind = shape or (MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE)
    x, y, w, h = int(x), int(y), int(w), int(h)
    box = s.shapes.add_shape(kind, x, y, w, h)
    if radius and kind == MSO_SHAPE.ROUNDED_RECTANGLE:
        box.adjustments[0] = min(0.5, radius / min(w, h))
    if fill is None:
        box.fill.background()
    else:
        box.fill.solid()
        box.fill.fore_color.rgb = fill
    if line is None:
        box.line.fill.background()
    else:
        box.line.color.rgb = line
        box.line.width = line_w
    box.shadow.inherit = False
    return box


def label(s, x, y, w, value, color=MUTED, size=10.5, align=PP_ALIGN.LEFT):
    """Жижиг том үсгэн гарчиг (eyebrow)."""
    return text(s, x, y, w, Cm(0.6), value.upper(), size=size, bold=True, color=color, spacing=1.2, align=align)


def shape_text(box, value, size=14, color=INK, bold=False, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE):
    frame = box.text_frame
    frame.word_wrap = True
    frame.margin_left = frame.margin_right = Cm(0.25)
    frame.margin_top = frame.margin_bottom = Cm(0.1)
    frame.vertical_anchor = anchor
    lines = value if isinstance(value, list) else [value]
    for i, item in enumerate(lines):
        p = frame.paragraphs[0] if i == 0 else frame.add_paragraph()
        p.alignment = align
        for part in (item if isinstance(item, list) else [item]):
            v, style = part if isinstance(part, tuple) else (part, {})
            run = p.add_run()
            run.text = v
            run.font.name = style.get('font', FONT)
            run.font.size = Pt(style.get('size', size))
            run.font.bold = style.get('bold', bold)
            run.font.color.rgb = style.get('color', color)
            if style.get('spacing'):
                run._r.get_or_add_rPr().set('spc', str(int(style['spacing'] * 100)))


def picture(s, name, x, y, w, h, border=None, valign='middle'):
    """Зургийг хайрцагт багтааж, харьцааг нь хадгалан голлуулна."""
    path = os.path.join(FIG, name)
    iw, ih = Image.open(path).size
    x, y, w, h = int(x), int(y), int(w), int(h)
    if iw / ih > w / h:
        pw, ph = w, int(w * ih / iw)
    else:
        pw, ph = int(h * iw / ih), h
    top = y if valign == 'top' else y + (h - ph) // 2
    pic = s.shapes.add_picture(path, x + (w - pw) // 2, top, pw, ph)
    if border is not None:
        pic.line.color.rgb = border
        pic.line.width = Pt(0.75)
    return pic


def phone(s, name, x, y, h):
    """Гар утасны дэлгэцийн зургийг хүрээтэй харуулна; өргөнийг буцаана."""
    path = os.path.join(FIG, name)
    iw, ih = Image.open(path).size
    pad = Cm(0.2)
    x, y, h = int(x), int(y), int(h)
    inner_h = h - 2 * pad
    inner_w = int(inner_h * iw / ih)
    rect(s, x, y, inner_w + 2 * pad, h, NAVY, radius=Cm(0.45))
    s.shapes.add_picture(path, x + pad, y + pad, inner_w, inner_h)
    return inner_w + 2 * pad


def code(s, x, y, w, h, source, size=11):
    """Псевдо код: бараан дэвсгэр, // тайлбарыг бүдэг өнгөөр."""
    rect(s, x, y, w, h, NAVY, radius=Cm(0.25))
    lines = []
    for raw in source.strip('\n').split('\n'):
        body, _, note = raw.partition('//')
        runs = [(body, {'color': CODE_INK})]
        if note:
            runs.append(('//' + note, {'color': CODE_NOTE}))
        lines.append(runs)
    return text(s, x + Cm(0.45), y + Cm(0.35), w - Cm(0.9), h - Cm(0.7), lines, size=size, font=MONO,
                after=0, line=1.0)


MARKS = {'y': ('✓', GREEN), 'p': ('∼', AMBER), 'n': ('–', GREY)}


def _border(cell, sides=('B',), color='D8DEE4', width=9525):
    tc_pr = cell._tc.get_or_add_tcPr()
    for side in reversed(('L', 'R', 'T', 'B')):
        if side not in sides:
            continue
        ln = etree.Element(qn('a:ln' + side), w=str(width), cap='flat', cmpd='sng', algn='ctr')
        fill = etree.SubElement(ln, qn('a:solidFill'))
        etree.SubElement(fill, qn('a:srgbClr'), val=color)
        etree.SubElement(ln, qn('a:prstDash'), val='solid')
        tc_pr.insert(0, ln)


def table(s, x, y, widths, rows, size=12, row_h=Cm(0.9), head=True, align=None, fills=None, bold_cols=(),
          head_fill=NAVY, head_color=WHITE, zebra=True, anchor=MSO_ANCHOR.MIDDLE, heights=None):
    """Хүснэгт: толгой мөр бараан, мөрүүд ээлжлэн цайвар, доод зураас нимгэн.

    Олон мөртэй нүдэнд heights-ээр өндрийг нь заана: PowerPoint мөрийг өөрөө сунгадаг ч Canva
    импортлохдоо хадгалсан өндрийг ашиглаж, илүү мөрийг тасалдаг.
    """
    heights = heights or [row_h] * len(rows)
    frame = s.shapes.add_table(len(rows), len(widths), x, y, sum(widths), sum(heights))
    tbl = frame.table
    pr = tbl._tbl.tblPr
    pr.set('firstRow', '0')
    pr.set('bandRow', '0')
    style = pr.find(qn('a:tableStyleId'))
    if style is None:
        style = etree.SubElement(pr, qn('a:tableStyleId'))
    style.text = '{2D5ABB26-0587-4C30-8999-92F81FD0307C}'  # No Style, No Grid
    for j, w in enumerate(widths):
        tbl.columns[j].width = w
    for i, row in enumerate(rows):
        tbl.rows[i].height = heights[i]
        is_head = head and i == 0
        for j, value in enumerate(row):
            cell = tbl.cell(i, j)
            cell.margin_left = cell.margin_right = Cm(0.18)
            cell.margin_top = cell.margin_bottom = Cm(0.06)
            cell.vertical_anchor = anchor
            fill = (fills or {}).get(j)
            if isinstance(fill, dict):
                fill = fill.get('head' if is_head else 'body')
            if fill is None:
                fill = head_fill if is_head else (CARD if zebra and i % 2 == 0 else WHITE)
            cell.fill.solid()
            cell.fill.fore_color.rgb = fill
            if not is_head:
                _border(cell)
            frame_text = cell.text_frame
            frame_text.word_wrap = True
            mark = MARKS.get(value) if isinstance(value, str) and not is_head else None
            paragraphs = [value] if mark else str(value).split('\n')
            for k, line_value in enumerate(paragraphs):
                p = frame_text.paragraphs[0] if k == 0 else frame_text.add_paragraph()
                p.alignment = (align or {}).get(j, PP_ALIGN.LEFT) if not mark else PP_ALIGN.CENTER
                p.space_after = Pt(1)
                run = p.add_run()
                if mark:
                    run.text, colour = mark
                    run.font.name = SYMBOL
                    run.font.size = Pt(size + 2)
                    run.font.bold = True
                    run.font.color.rgb = colour
                    continue
                run.text = line_value
                run.font.name = FONT
                run.font.size = Pt(size)
                run.font.bold = is_head or j in bold_cols
                run.font.color.rgb = head_color if is_head else INK
    return tbl


SECTIONS = {
    1: '01 · Хэрэгцээ ба зорилго',
    2: '02 · Онол, аргазүй, шинжилгээ',
    3: '03 · Шаардлага ба use-case',
    4: '04 · Зохиомж ба төлөвлөгөө',
    0: 'Нэмэлт слайд',
}


def slide(title, section, notes, number=None):
    s = prs.slides.add_slide(blank)
    label(s, M, Cm(0.85), Cm(24), SECTIONS[section])
    rect(s, M, Cm(1.5), Cm(1.1), Cm(0.13), YELLOW)
    text(s, M, Cm(1.72), Cm(28), Cm(1.4), title, size=28, bold=True, color=NAVY)
    n = number or '%02d' % len(prs.slides._sldIdLst)
    text(s, W - M - Cm(3), Cm(0.8), Cm(3), Cm(0.7), n, size=13, bold=True, color=NAVY, align=PP_ALIGN.RIGHT)
    rect(s, M, Cm(18.05), CW, Emu(9525), LINE)
    text(s, M, Cm(18.2), Cm(22), Cm(0.5), AUTHOR + ' · ' + TITLE, size=9, color=MUTED)
    text(s, W - M - Cm(10), Cm(18.2), Cm(10), Cm(0.5), 'Дипломын ажлын I үзлэг · 2026.09.30', size=9,
         color=MUTED, align=PP_ALIGN.RIGHT)
    s.notes_slide.notes_text_frame.text = notes
    return s


def dark_slide(notes):
    s = prs.slides.add_slide(blank)
    rect(s, 0, 0, W, H, NAVY)
    s.notes_slide.notes_text_frame.text = notes
    return s


# --------------------------------------------------------------------------- 1. Нүүр
s = dark_slide('Сайн байна уу. Би Программ хангамжийн инженерчлэлийн оюутан Б.Билгүүн. Дипломын сэдэв маань '
               '«Байгууллагын бүтээмжийг дэмжих платформ хөгжүүлэлт». Энэ бол 5S аудит, ажлын удирдлага, lean '
               'дадал, тайланг нэгтгэсэн, өөрийн санаачилгаар хөгжүүлж буй вэб ба гар утасны систем. Удирдагч '
               'багш маань доктор (Ph.D), дэд профессор Г.Ганчимэг.')
rect(s, 0, 0, Cm(0.45), H, YELLOW)
rect(s, Cm(21.2), Cm(12.2), W - Cm(21.2), H - Cm(12.2), NAVY2)
label(s, Cm(1.9), Cm(1.5), Cm(20), 'ШУТИС · МХТС · Компьютерийн ухааны тэнхим', color=YELLOW, size=11)
text(s, Cm(1.9), Cm(2.25), Cm(18), Cm(0.7), 'Бакалаврын дипломын ажил · I үзлэг', size=13, color=PALE)
text(s, Cm(1.9), Cm(4.0), Cm(18), Cm(5.4), ['Байгууллагын', 'бүтээмжийг дэмжих', 'платформ хөгжүүлэлт'],
     size=40, bold=True, color=WHITE, after=0, line=1.0)
text(s, Cm(1.9), Cm(9.35), Cm(17.5), Cm(0.8), 'Development of a Platform to Support Organizational Productivity',
     size=13, color=PALE, italic=True)
rect(s, Cm(1.9), Cm(12.3), Cm(1.1), Cm(0.13), YELLOW)
text(s, Cm(1.9), Cm(12.75), Cm(18.5), Cm(4.5), [
    [('Гүйцэтгэсэн   ', {'color': PALE}), ('Б.Билгүүн (B190910014)', {'bold': True})],
    [('Удирдагч   ', {'color': PALE}), ('Доктор (Ph.D), дэд профессор Г.Ганчимэг', {'bold': True})],
    [('Мэргэжил   ', {'color': PALE}), ('Программ хангамжийн инженерчлэл · 071405000000002306', {})],
    [('Огноо   ', {'color': PALE}), ('2026.09.30 · Улаанбаатар', {})],
], size=13, color=WHITE, after=7)
picture(s, 'ui_huddle.png', Cm(19.2), Cm(2.9), Cm(13.3), Cm(8.4), border=RGBColor(0x2B, 0x3A, 0x4A), valign='top')
phone(s, 'mob_fives.png', Cm(27.9), Cm(7.9), Cm(9.9))

# --------------------------------------------------------------------------- 2. Агуулга
s = slide('Агуулга', 1, 'Илтгэл дөрвөн хэсэгтэй: хэрэгцээ ба зорилго, онол ба шинжилгээ, шаардлага ба use-case, '
          'зохиомж ба төлөвлөгөө. Баруун талд нэгдүгээр үзлэгийн шаардлага бүр аль слайдад байгааг харууллаа.')
parts = [('01', 'Хэрэгцээ ба зорилго', 'Шийдэх асуудал · зорилго · зорилт'),
         ('02', 'Онол, аргазүй, шинжилгээ', '5S, Gemba, Kaizen · PDCA · бүтээмжийн KPI · ижил төстэй систем · SWOT'),
         ('03', 'Шаардлага ба use-case', '22 функциональ, 10 функциональ бус шаардлага · 9 use-case · мөрдөх матриц'),
         ('04', 'Зохиомж ба төлөвлөгөө', 'Архитектур · өгөгдлийн сан · алгоритм · интерфейс · төлөв · ном зүй')]
for i, (num, head, sub) in enumerate(parts):
    y = Cm(3.9 + i * 3.4)
    badge = rect(s, M, y, Cm(1.6), Cm(1.6), YELLOW if i == 0 else NAVY)
    shape_text(badge, num, size=18, bold=True, color=NAVY if i == 0 else WHITE)
    text(s, M + Cm(2.2), y + Cm(0.05), Cm(15.6), Cm(0.9), head, size=20, bold=True, color=INK)
    text(s, M + Cm(2.2), y + Cm(0.95), Cm(15.6), Cm(1.4), sub, size=12.5, color=MUTED)
card_x, card_w = Cm(20.3), W - M - Cm(20.3)
rect(s, card_x, Cm(3.9), card_w, Cm(13.5), NAVY, radius=Cm(0.3))
label(s, card_x + Cm(0.8), Cm(4.55), card_w - Cm(1.6), 'I үзлэгийн шаардлага', color=YELLOW)
checks = [('Сэдвийн хэрэгцээ, шийдэх асуудал', '3'), ('Онол, аргазүйн судалгаа', '5–7'),
          ('Ижил төстэй систем, SWOT', '8–9'), ('Функцийн шаардлага', '10'),
          ('Use-case диаграм, тодорхойлолт', '11–13'), ('Архитектур, загвар', '14–16'),
          ('Алгоритмын зохиомж', '17–18'), ('Ном зүй (31 эх сурвалж)', '21')]
for i, (item, pages) in enumerate(checks):
    y = Cm(5.55 + i * 1.44)
    text(s, card_x + Cm(0.8), y, Cm(0.8), Cm(0.8), '✓', size=15, bold=True, color=YELLOW, font=SYMBOL)
    text(s, card_x + Cm(1.6), y + Cm(0.03), card_w - Cm(3.8), Cm(0.8), item, size=13.5, color=WHITE)
    text(s, card_x + card_w - Cm(2.2), y + Cm(0.03), Cm(1.4), Cm(0.8), pages, size=13, color=PALE,
         align=PP_ALIGN.RIGHT)

# --------------------------------------------------------------------------- 3. Хэрэгцээ, асуудал
s = slide('Сэдвийн хэрэгцээ ба шийдэх асуудал', 1,
          '5S нэвтрүүлсэн байгууллагад аудит цаасаар хийгдэж, илэрсэн дутагдлыг хэн, хэзээ засах нь хянагддаггүй. '
          'Ажил ам, чатаар өгөгдөж, хоцролт хурал дээр л илэрдэг. Сарын тайланг гараар нэгтгэдэг. Санаа, түүх '
          'бүртгэгддэггүй. Иймээс эдгээрийг нэг өгөгдлийн санд нэгтгэсэн, монгол хэлтэй, сүлжээгүй орчинд '
          'ажилладаг платформ хэрэгтэй гэж үзсэн.')
label(s, M, Cm(3.65), Cm(10), 'Одоогийн байдал')
label(s, Cm(14.8), Cm(3.65), Cm(10), 'Үр дагавар')
problems = [('5S аудит цаасаар хийгдэж, дараа нь Excel-д шивэгдэнэ',
             'Илэрсэн дутагдлыг хэн, хэзээ засахыг хянадаггүй — аудит хэмжилт төдий үлддэг'),
            ('Ажил ам, чатаар өгөгдөнө',
             'Менежер явц, хоцролтыг бодит хугацаанд хардаггүй; хоцролт хурал дээр л илэрдэг'),
            ('Сарын тайланг хэлтэс бүрээс гараар нэгтгэнэ',
             'Сар бүр олон цаг зарцуулж, алдаа гаргах эрсдэлтэй'),
            ('Ажилтны санаа, шийдвэрийн түүх бүртгэгддэггүй',
             'Санаа гаргагч хариугаа авдаггүй, байгууллагын туршлага алдагддаг')]
for i, (now, effect) in enumerate(problems):
    y = Cm(4.35 + i * 2.55)
    box = rect(s, M, y, Cm(12.2), Cm(2.15), CARD, radius=Cm(0.2))
    shape_text(box, now, size=15, bold=True, color=INK, align=PP_ALIGN.LEFT)
    arrow = rect(s, Cm(14.0), y + Cm(0.72), Cm(0.6), Cm(0.7), YELLOW, shape=MSO_SHAPE.RIGHT_ARROW)
    box = rect(s, Cm(14.8), y, W - M - Cm(14.8), Cm(2.15), WHITE, radius=Cm(0.2), line=LINE)
    shape_text(box, effect, size=14, color=INK, align=PP_ALIGN.LEFT)
box = rect(s, M, Cm(14.75), CW, Cm(2.75), NAVY, radius=Cm(0.25))
shape_text(box, [[('ШИЙДЭЛ   ', {'color': YELLOW, 'size': 12, 'spacing': 1.2}),
                  ('5S аудит, ажил, lean дадал, тайланг нэг өгөгдлийн санд нэгтгэсэн, монгол хэлтэй, сүлжээгүй '
                   'орчинд ажилладаг вэб ба гар утасны платформ', {})]],
           size=17, bold=True, color=WHITE, align=PP_ALIGN.LEFT)

# --------------------------------------------------------------------------- 4. Зорилго, зорилт
s = slide('Зорилго ба зорилт', 1,
          'Зорилго: 5S, бүтээмжийн хэрэгсэл, ажлын удирдлагыг нэг платформд нэгтгэх. Зорилтыг гурван үе шатад '
          'хуваасан: судлах, зохиомжлох ба хөгжүүлэх, үнэлэх. Судлах, зохиомжлох хэсэг энэ үзлэгийн агуулга; '
          'туршилт, үнэлгээ III үзлэгт.')
rect(s, M, Cm(3.75), Cm(0.16), Cm(2.75), YELLOW)
label(s, M + Cm(0.6), Cm(3.65), Cm(10), 'Зорилго')
text(s, M + Cm(0.6), Cm(4.3), CW - Cm(0.6), Cm(2.4),
     'Байгууллагын 5S, бүтээмжийн хэрэгслүүд болон ажлын удирдлагыг нэгтгэж, менежерт ажлыг бодит хугацаанд '
     'оноож хянах, ажилтанд ажлаа гар утаснаас хөтлөх, тайланг автоматаар гаргах боломж олгох вэб ба гар утасны '
     'платформыг зохион бүтээж хэрэгжүүлэх', size=17, bold=True, color=NAVY, line=1.12)
goals = [('01', 'Судлах', 'I үзлэг', ['5S, Gemba, Kaizen, PDCA, шатлалт хурлын онолыг судлах',
                                      'Бүтээмжийн үзүүлэлт (KPI), үнэлгээний загварыг тодорхойлох',
                                      'Ижил төстэй системийг харьцуулан шинжлэх', 'Хэрэглэгчийн судалгаа авах']),
         ('02', 'Зохиомжлох, хөгжүүлэх', 'I–II үзлэг', ['Функциональ, функциональ бус шаардлага',
                                                       'UML, архитектур, өгөгдлийн сан, алгоритмын зохиомж',
                                                       'Backend API, вэб, гар утасны апп, автомат тест']),
         ('03', 'Үнэлэх', 'III үзлэг', ['Туршилтын байгууллагад ажиллуулах', 'KPI-г «өмнө–дараа» загвараар хэмжих',
                                       'SUS аргаар хэрэглэхэд хялбар байдлыг үнэлэх', 'Эдийн засгийн үр ашиг тооцох'])]
col_w = (CW - Cm(1.2)) / 3
for i, (num, head, tag, items) in enumerate(goals):
    x = M + int(i * (col_w + Cm(0.6)))
    rect(s, x, Cm(7.35), int(col_w), Cm(10.15), CARD, radius=Cm(0.25))
    badge = rect(s, x + Cm(0.6), Cm(7.95), Cm(1.1), Cm(1.1), NAVY)
    shape_text(badge, num, size=13, bold=True, color=WHITE)
    text(s, x + Cm(2.0), Cm(8.1), int(col_w) - Cm(2.4), Cm(0.9), head, size=16, bold=True, color=INK)
    chip = rect(s, x + Cm(2.0), Cm(9.05), Cm(2.5), Cm(0.72), WHITE, radius=Cm(0.36), line=LINE)
    shape_text(chip, tag, size=10.5, bold=True, color=MUTED)
    text(s, x + Cm(0.6), Cm(10.4), int(col_w) - Cm(1.2), Cm(6.8), ['•  ' + t for t in items], size=14,
         color=INK, after=9)

# --------------------------------------------------------------------------- 5. Онол
s = slide('Онол, аргазүйн судалгаа', 2,
          'Онолын суурь нь дөрвөн аргачлал. 5S — ажлын байрны эмх цэгц, түүнийг аудитаар хэмждэг. Gemba, Kaizen — '
          'удирдлага ажлын байранд очиж, жижиг сайжруулалтыг тасралтгүй хийх. Шатлалт өдөр тутмын хурал — ээлжийн '
          'эхний богино хурал, харааны самбар. PDCA ба KPI — сайжруулалтын мөчлөг ба хэмжилт. Карт бүрийн доод мөр '
          'онол системийн аль хэсэгт хэрэгжсэнийг заана.')
theory = [('5S', 'Hirano, 1995; Osada, 1991',
           ['Ангилах · Цэгцлэх · Цэвэрлэх · Стандартчлах · Төлөвшүүлэх',
            'Түвшинг шалгах хуудсаар аудит хийж, оноогоор хэмжинэ'],
           'Талбай, бүс, QR аудит, улаан шошго, аудитаас засах ажил'),
          ('Gemba ба Kaizen', 'Imai, 1997; Liker, 2004',
           ['Удирдлага ажлын байранд очиж ажиглана',
            'Жижиг, тасралтгүй сайжруулалт, ажилтны санаа'],
           'Gemba явалт → ажил, сайжруулалтын санаа ба хариу'),
          ('Шатлалт өдөр тутмын хурал', 'Mann, 2005',
           ['Ээлжийн эхэнд 10–15 минутын хурал, харааны самбар',
            'Удирдагчийн стандарт ажил, долоо хоногийн тайлан'],
           'Өглөөний хурлын самбар, долоо хоногийн тайлан, сануулга'),
          ('PDCA ба бүтээмжийн хэмжилт', 'Deming, 1986; Parmenter, 2015; OECD, 2001',
           ['Төлөвлөх → Хийх → Шалгах → Сайжруулах мөчлөг',
            'Бүтээмж = гарц / орц; байгууллагын түвшинд процессын KPI-ээр'],
           'K1–K7 үзүүлэлт, автомат сарын тайлан')]
card_w = (CW - Cm(0.6)) / 2
for i, (head, source, points, system) in enumerate(theory):
    x = M + int((i % 2) * (card_w + Cm(0.6)))
    y = Cm(3.75 + (i // 2) * 7.0)
    rect(s, x, y, int(card_w), Cm(6.6), CARD, radius=Cm(0.25))
    text(s, x + Cm(0.7), y + Cm(0.45), int(card_w) - Cm(1.4), Cm(0.9), head, size=19, bold=True, color=NAVY)
    text(s, x + Cm(0.7), y + Cm(1.3), int(card_w) - Cm(1.4), Cm(0.6), source, size=10.5, color=MUTED,
         italic=True)
    text(s, x + Cm(0.7), y + Cm(1.95), int(card_w) - Cm(1.4), Cm(2.8), ['•  ' + p for p in points], size=14,
         after=4)
    rect(s, x + Cm(0.7), y + Cm(4.95), int(card_w) - Cm(1.4), Emu(9525), LINE)
    text(s, x + Cm(0.7), y + Cm(5.2), int(card_w) - Cm(1.4), Cm(1.2),
         [[('СИСТЕМД  ', {'color': GREEN, 'bold': True, 'size': 10.5, 'spacing': 1.2}), (system, {})]],
         size=13.5, color=INK)

# --------------------------------------------------------------------------- 6. PDCA
s = slide('PDCA мөчлөг ба системийн холбоос', 2,
          'PDCA мөчлөгийн алхам бүрт системийн функц харгалзана. Хамгийн сул холбоос бол «Сайжруулах» алхмын үр дүнг '
          'дахин хэмжих явдал: засах ажил хаагдсан ч асуудал арилсан эсэх тодорхойгүй үлддэг. Системд аудитын оноо '
          'босгоос доош бол засах ажил автоматаар үүсч, дараагийн аудитаар үр дүнг шалгана; засах хугацааг K4 '
          'үзүүлэлтээр хэмжинэ.')
picture(s, 'pdca.png', M, Cm(3.4), Cm(19.2), Cm(14.3))
x = Cm(21.4)
w = W - M - x
rect(s, x, Cm(3.9), w, Cm(13.6), CARD, radius=Cm(0.3))
label(s, x + Cm(0.8), Cm(4.6), w - Cm(1.6), 'Сул холбоос', color=RED)
text(s, x + Cm(0.8), Cm(5.3), w - Cm(1.6), Cm(3.2),
     '«Сайжруулах» алхмын үр дүнг дахин хэмжих: засах ажил хаагдсан ч асуудал арилсан эсэх тодорхойгүй үлддэг.',
     size=15, bold=True, color=INK, line=1.12)
label(s, x + Cm(0.8), Cm(8.85), w - Cm(1.6), 'Системийн шийдэл', color=GREEN)
text(s, x + Cm(0.8), Cm(9.55), w - Cm(1.6), Cm(6.0),
     ['•  Аудитын оноо < 85% → бүсийн хариуцагчид засах ажил автоматаар',
      '•  Үр дүнг дараагийн аудитаар шалгана',
      '•  Засах ажлын хаалтын хугацааг K4-өөр хэмжинэ',
      '•  Өглөөний хурал, сарын тайлан «Шалгах» алхмыг хөтөлнө'], size=14, after=8)
text(s, x + Cm(0.8), Cm(16.45), w - Cm(1.6), Cm(0.6), 'Эх сурвалж: Deming (1986)', size=10.5, color=MUTED,
     italic=True)

# --------------------------------------------------------------------------- 7. KPI
s = slide('Бүтээмжийг хэмжих аргачлал', 2,
          'Бүтээмжийг долоон үзүүлэлтээр хэмжинэ. K1–K4, K6-ийн эх өгөгдөл системд аль хэдийн бүртгэгддэг. '
          'Үнэлгээг «өмнө–дараа» загвараар хийнэ: хоёр долоо хоногийн суурь хэмжилт, дөрвөн долоо хоногийн туршилт. '
          'Нэг байгууллага, богино хугацаа тул үр дүнг учир шалтгааны баталгаа бус, чиг хандлага гэж тайлбарлана. '
          'SUS-ийн 68 оноо нь олон судалгааны дундаж тул жишиг болгосон.')
kpi = [['Код', 'Үзүүлэлт', 'Тооцох арга', 'Эх өгөгдөл', 'Чиглэл'],
       ['K1', 'Хугацаандаа гүйцэтгэл', 'Хугацаандаа дууссан / хугацаа нь болсон ажил × 100%', 'tasks', '↑ өсөх'],
       ['K2', 'Хоцорсон ажил', 'Хугацаа хэтэрсэн ажлын тоо, дундаж хоцролт (хоног)', 'tasks', '↓ буурах'],
       ['K3', '5S-ийн түвшин', 'Аудитын дундаж оноо; ≥ 85% хүрсэн бүсийн хувь', 'audit_runs', '↑ өсөх'],
       ['K4', 'Засах ажлын хаалт', 'Аудитаас үүссэн ажлын дундаж хаалт (хоног)', 'tasks (аудит)', '↓ буурах'],
       ['K5', 'Тайлан бэлтгэх хугацаа', 'Сарын тайланд зарцуулсан цаг: өмнө ба дараа', 'менежерийн тэмдэглэл',
        '↓ буурах'],
       ['K6', 'Хэрэглээ', 'Долоо хоногийн тайлангаа бичсэн ажилтны хувь', 'weekly_checkins', '↑ өсөх'],
       ['K7', 'Хэрэглэхэд хялбар байдал', 'SUS асуулгын дундаж оноо (0–100)', 'асуулга', '≥ 68']]
table(s, M, Cm(3.7), [Cm(1.5), Cm(6.4), Cm(13.4), Cm(5.4), Cm(3.967)], kpi, size=12.5, row_h=Cm(0.98),
      bold_cols=(0,), align={0: PP_ALIGN.CENTER, 4: PP_ALIGN.CENTER})
label(s, M, Cm(12.05), Cm(20), 'Үнэлгээний загвар: өмнө – дараа')
steps = [('Суурь хэмжилт · 2 долоо хоног', 'Бүртгэл, ярилцлагаас суурь утга', CARD, INK),
         ('Туршилт · 4 долоо хоног', 'K1–K6 долоо хоног бүр', YELLOW, NAVY),
         ('Дүгнэлт', 'K7 (SUS), K5; өөрчлөлтийг чиг хандлагаар', NAVY, WHITE)]
step_w = (CW - Cm(0.3)) / 3
for i, (head, sub, fill, ink) in enumerate(steps):
    x = M + int(i * (step_w + Cm(0.15)))
    box = rect(s, x, Cm(12.75), int(step_w), Cm(2.55), fill,
               shape=MSO_SHAPE.PENTAGON if i == 0 else MSO_SHAPE.CHEVRON)
    box.adjustments[0] = 0.28
    shape_text(box, [[(head, {'bold': True, 'size': 14})], [(sub, {'size': 12})]], color=ink)
text(s, M, Cm(15.85), CW, Cm(1.6),
     'Нэг байгууллага, богино хугацаатай туршилт тул үр дүнг учир шалтгааны баталгаа бус, чиг хандлага гэж '
     'тайлбарлана. Эх сурвалж: Parmenter (2015); Brooke (1996); Bangor ба бусад (2008).', size=11.5, color=MUTED)

# --------------------------------------------------------------------------- 8. Харьцуулалт
s = slide('Ижил төстэй системийн шинжилгээ', 2,
          'Таван ижил төстэй системийг арван боломжоор харьцуулсан. SafetyCulture аудит, сүлжээгүй горимоороо, '
          'Weekdone долоо хоногийн тайлангаараа, KaiNexus ба Tervene сайжруулалт, lean дадлаараа сайн. Гэвч бүгдийг '
          'нэг дор, монгол хэлээр, өөрийн серверт байршуулах шийдэл алга; хэрэглэгч тус бүрээр валютаар төлбөртэй.')
cmp_rows = [['Боломж', 'SafetyCulture', 'Weekdone', 'Asana/Trello', 'KaiNexus', 'Tervene', 'Энэ систем'],
            ['Аудит, шалгах хуудас (утаснаас)', 'y', 'n', 'n', 'p', 'y', 'y'],
            ['Ажил оноох, хянах', 'p', 'p', 'y', 'y', 'y', 'y'],
            ['Аудитаас автомат засах ажил', 'y', 'n', 'n', 'n', 'p', 'y'],
            ['Сайжруулалтын санаа', 'n', 'n', 'n', 'y', 'y', 'y'],
            ['Өдөр тутмын хурлын самбар', 'n', 'n', 'n', 'p', 'y', 'y'],
            ['Долоо хоногийн тайлан', 'n', 'y', 'n', 'n', 'n', 'y'],
            ['Gemba явалтын бүртгэл', 'p', 'n', 'n', 'n', 'y', 'y'],
            ['Сарын тайлан автоматаар', 'p', 'p', 'n', 'p', 'p', 'y'],
            ['Сүлжээгүй горим', 'y', 'n', 'p', 'n', 'n', 'y'],
            ['Монгол хэл, өөрийн серверт', 'n', 'n', 'n', 'n', 'n', 'y']]
widths = [Cm(9.2)] + [int((CW - Cm(9.2)) / 6)] * 6
cmp_table = table(s, M, Cm(3.65), widths, cmp_rows, size=12.5, row_h=Cm(0.93),
                  align={j: PP_ALIGN.CENTER for j in range(1, 7)}, fills={6: {'head': YELLOW, 'body': TINT}})
cmp_table.cell(0, 6).text_frame.paragraphs[0].runs[0].font.color.rgb = NAVY
text(s, M, Cm(14.05), CW, Cm(0.6), [[('✓', {'font': SYMBOL, 'color': GREEN, 'bold': True}), (' бүрэн     ', {}),
                                     ('∼', {'font': SYMBOL, 'color': AMBER, 'bold': True}), (' хэсэгчлэн     ', {}),
                                     ('–', {'color': GREY, 'bold': True}), (' байхгүй', {})]],
     size=11, color=MUTED)
box = rect(s, M, Cm(14.85), CW, Cm(2.65), NAVY, radius=Cm(0.25))
shape_text(box, [[('ДҮГНЭЛТ   ', {'color': YELLOW, 'size': 12, 'bold': True}),
                  ('Тус бүр нэг хэсэгтээ сайн ч аудит, ажил, lean дадлыг нэг дор, монгол хэлээр, өөрийн серверт '
                   'байршуулдаг шийдэл алга. Шилдэг туршлагыг (офлайн аудит, хариуцагчтай засах ажил, хурлын самбар, '
                   'долоо хоногийн сануулга) нэгтгэв.', {})]], size=14, color=WHITE, align=PP_ALIGN.LEFT)

# --------------------------------------------------------------------------- 9. SWOT
s = slide('SWOT шинжилгээ', 2,
          'SWOT: давуу тал — нэгдсэн систем, аудитаас автомат засах ажил, офлайн горим, өөрийн серверт байршуулах. '
          'Сул тал — бодит хэрэглэгчийн туршлага бага, push мэдэгдэл байхгүй, нэг хөгжүүлэгч. Боломж — 5S нэвтрүүлж '
          'буй байгууллага олон, тайлангийн цаг хэмнэх хэрэгцээ. Аюул — гадаадын системүүдтэй өрсөлдөх, цахим '
          'шилжилтийн эсэргүүцэл, хувийн мэдээллийн эрсдэл.')
swot = [('S', 'Давуу тал', GREEN, ['Аудит, ажил, lean дадал нэг системд', 'Аудитаас засах ажил автоматаар үүснэ',
                                   'Сүлжээгүй горим, монгол хэл', 'Өөрийн серверт байршуулах боломж']),
        ('W', 'Сул тал', NAVY, ['Хэрэглэгчийн бодит туршлага бага', 'Push мэдэгдэл хараахан байхгүй',
                                'Нэг хөгжүүлэгчтэй, дэмжлэгийн нөөц хязгаарлагдмал']),
        ('O', 'Боломж', BLUE, ['5S, lean нэвтрүүлж буй байгууллага олон', 'Тайлан бэлтгэх цаг хэмнэх хэрэгцээ өндөр',
                               'AI ашиглан тайлан ноороглох']),
        ('T', 'Аюул', RED, ['Гадаадын том системүүдтэй өрсөлдөх', 'Байгууллагын цахим шилжилтэд эсэргүүцэл',
                            'Хувийн мэдээллийн аюулгүй байдлын эрсдэл'])]
for i, (letter, head, colour, items) in enumerate(swot):
    x = M + int((i % 2) * (card_w + Cm(0.6)))
    y = Cm(3.75 + (i // 2) * 7.0)
    rect(s, x, y, int(card_w), Cm(6.6), CARD, radius=Cm(0.25))
    badge = rect(s, x + Cm(0.7), y + Cm(0.6), Cm(1.5), Cm(1.5), colour, radius=Cm(0.2))
    shape_text(badge, letter, size=24, bold=True, color=WHITE)
    text(s, x + Cm(2.6), y + Cm(0.85), int(card_w) - Cm(3.2), Cm(1.0), head, size=19, bold=True, color=INK)
    text(s, x + Cm(0.7), y + Cm(2.55), int(card_w) - Cm(1.4), Cm(3.9), ['•  ' + t for t in items], size=14.5,
         after=6)

# --------------------------------------------------------------------------- 10. Шаардлага
s = slide('Функциональ ба функциональ бус шаардлага', 3,
          '22 функциональ шаардлагыг зургаан модульд бүлэглэж, MoSCoW аргаар эрэмбэлсэн: 13 заавал, 8 байх ёстой, '
          '1 байж болно. Функциональ бус 10 шаардлага аюулгүй байдал, хүртээмж, найдвартай байдал, гүйцэтгэл, '
          'засварлагдах чанарыг хамарна. Дүрүүд үүрлэсэн: дээд дүр доод дүрийн эрхийг өвлөнө.')
label(s, M, Cm(3.65), Cm(12), 'Функциональ шаардлага · 22')
text(s, Cm(12.4), Cm(3.6), Cm(8.7), Cm(0.6), [[('MoSCoW:  ', {'color': MUTED}), ('M 13', {'bold': True}),
                                              ('  ·  ', {'color': GREY}), ('S 8', {'bold': True}),
                                              ('  ·  ', {'color': GREY}), ('C 1', {'bold': True})]],
     size=11.5, color=INK, align=PP_ALIGN.RIGHT)
modules = [('Нэвтрэлт ба эрх', [('FR-01', 'Нэвтрэх, эрхийн хүрээ', 'M'), ('FR-02', 'Урилга, эрх олгох', 'M'),
                                ('FR-21', 'Бүртгэл устгах хүсэлт', 'S')]),
           ('Ажлын удирдлага', [('FR-03', 'Ажил үүсгэх, оноох', 'M'), ('FR-04', 'Ач холбогдлоор эрэмбэлэх', 'M'),
                                ('FR-05', 'Оноолт, өглөөний мэдэгдэл', 'M')]),
           ('5S ба аудит', [('FR-06', 'Талбай, бүс, стандарт', 'M'), ('FR-07', 'QR-ээр аудит хийх', 'M'),
                            ('FR-08', 'Аудитаас засах ажил', 'M'), ('FR-09', 'Улаан шошго, зураг', 'M')]),
           ('Бүртгэл ба тайлан', [('FR-10', 'Өдрийн бүртгэл', 'M'), ('FR-11', 'Долоо хоногийн тайлан', 'S'),
                                  ('FR-12', 'Сарын тайлан, хаалт', 'M'), ('FR-13', 'Хагас жил, жилийн тайлан', 'S'),
                                  ('FR-14', 'AI тойм бичвэр', 'C'), ('FR-20', 'Үйл ажиллагааны түүх', 'S')]),
           ('Lean дадал', [('FR-15', 'Өглөөний хурлын самбар', 'S'), ('FR-16', 'Gemba явалт', 'S'),
                           ('FR-17', 'Сайжруулалтын санаа', 'S')]),
           ('Систем', [('FR-18', 'Сүлжээгүй горим', 'M'), ('FR-19', 'Монгол, англи хэл', 'M'),
                       ('FR-22', 'KPI самбар (K1–K6)', 'S')])]
mod_w = (Cm(19.5) - Cm(0.5)) / 2
for i, (head, items) in enumerate(modules):
    x = M + int((i % 2) * (mod_w + Cm(0.5)))
    y = Cm(4.4 + (i // 2) * 4.45)
    rect(s, x, y, int(mod_w), Cm(4.15), CARD, radius=Cm(0.2))
    text(s, x + Cm(0.5), y + Cm(0.3), int(mod_w) - Cm(1.0), Cm(0.7), head, size=14, bold=True, color=NAVY)
    text(s, x + Cm(0.5), y + Cm(0.95), int(mod_w) - Cm(1.0), Cm(3.1),
         [[(code_id + '  ', {'bold': True, 'color': MUTED, 'size': 10.5}), (name, {}),
           ('  ' + prio, {'color': GREY, 'size': 10})] for code_id, name, prio in items],
         size=11.5, after=1, line=1.0)
x = Cm(21.6)
w = W - M - x
rect(s, x, Cm(3.65), w, Cm(13.85), NAVY, radius=Cm(0.3))
label(s, x + Cm(0.7), Cm(4.25), w - Cm(1.4), 'Функциональ бус · 10', color=YELLOW)
nfr = [('01', 'Нууц үг hash, JWT, сервер талын эрх'), ('02', 'Байгууллагын өгөгдөл тусгаарлагдана'),
       ('03', 'WCAG 2.1 AA; SUS ≥ 68'), ('04', 'Нэг гараар, том үсэг, бараан горим'),
       ('05', 'Сүлжээ тасрахад өгөгдөл алдагдахгүй'), ('06', 'Гол хуудас < 2 секунд'),
       ('07', 'Migration, автомат тест'), ('08', 'Docker, аль ч Linux сервер'),
       ('09', 'AI-д зөвхөн нэргүй, нэгдсэн тоо'), ('10', 'Asia/Ulaanbaatar цагийн бүс')]
text(s, x + Cm(0.7), Cm(5.0), w - Cm(1.4), Cm(10.0),
     [[('NFR-' + n + '  ', {'color': YELLOW, 'bold': True, 'size': 11}), (t, {})] for n, t in nfr],
     size=12.5, color=WHITE, after=5.5, line=1.0)
text(s, x + Cm(0.7), Cm(15.75), w - Cm(1.4), Cm(1.5),
     'Дүрүүд үүрлэсэн: ажиглагч → ажилтан → менежер → админ', size=11, color=PALE)

# --------------------------------------------------------------------------- 11. Use-case диаграм
s = slide('Use-case диаграм', 3,
          'Таван оролцогч: ажиглагч, ажилтан, менежер, админ, мөн цагийн хуваарь буюу систем өөрөө. Ажилтан '
          'ажиглагчийн, админ менежерийн бүх use-case-ийг өвлөнө. Системийн хил дотор 18 use-case байгаагаас '
          'тодоор тэмдэглэсэн есийг нь (UC-01…UC-09) тоглогч, нөхцөл, урсгалтай нь дэлгэрэнгүй тодорхойлсон.')
picture(s, 'usecase_slide.png', M, Cm(3.5), Cm(25.2), Cm(14.0), valign='top')
x = Cm(27.3)
w = W - M - x
facts = [('5', 'оролцогч', 'ажиглагч, ажилтан, менежер, админ, цагийн хуваарь'),
         ('18', 'use-case', 'системийн хил доторх үйлдлүүд'),
         ('9', 'тодорхойлсон', 'UC-01…UC-09: тоглогч, нөхцөл, үндсэн ба алтернатив урсгал')]
for i, (num, name, note) in enumerate(facts):
    y = Cm(3.6 + i * 4.7)
    dark = i == 2
    rect(s, x, y, w, Cm(4.3), NAVY if dark else CARD, radius=Cm(0.25))
    text(s, x + Cm(0.5), y + Cm(0.3), w - Cm(1.0), Cm(1.3), num, size=30, bold=True, color=YELLOW if dark else NAVY)
    text(s, x + Cm(0.5), y + Cm(1.65), w - Cm(1.0), Cm(0.7), name, size=13, bold=True, color=WHITE if dark else INK)
    text(s, x + Cm(0.5), y + Cm(2.35), w - Cm(1.0), Cm(1.8), note, size=10.5, color=PALE if dark else MUTED)

# --------------------------------------------------------------------------- 12. Use-case тодорхойлолт
s = slide('Use-case тодорхойлолт: UC-02 «5S аудит хийх»', 3,
          'Жишээ болгон 5S аудит хийх use-case-ийг харууллаа. Ажилтан QR шошго уншуулж, шалгах хуудсаар оноо өгнө; '
          'оноо 85%-иас доош бол систем бүсийн хариуцагчид засах ажил үүсгэж мэдэгдэнэ. Сүлжээгүй үед утсанд '
          'хадгалагдаж, дараа илгээгдэнэ; нэг аудитаас ажил давхар үүсэхгүй. Баруун талд энэ use-case аль '
          'шаардлага, алгоритм, диаграмм, тесттэй холбогдохыг харууллаа; дарааллын диаграмм нь нэмэлт слайд Н2-т. '
          'Бусад найман use-case-ийг дипломын 2-р бүлэгт ижил загвараар тодорхойлсон.')
uc_rows = [['Юзкейс', '5S аудит хийх · UC-02'],
           ['Тоглогч', 'Үндсэн: ажилтан · Нэмэлт: систем'],
           ['Товч тайлбар', 'Бүсийн 5S түвшинг шалгах хуудсаар хэмжиж, дутагдлыг бүртгэнэ'],
           ['Өмнөх нөхцөл', 'Нэвтэрсэн; бүс аудитын загвартай; audits:create эрхтэй'],
           ['Үндсэн урсгал', '1. Ажилтан бүсийн QR шошгыг уншуулна\n'
                             '2. Систем шалгах хуудас, стандарт зургийг харуулна\n'
                             '3. Асуулт бүрт оноо, тийм/үгүй хариулт, шаардлагатай бол зураг\n'
                             '4. Ажилтан «Дуусгах» дарна\n'
                             '5. Систем оноог тооцож хадгална, бүсийн оноог шинэчилнэ\n'
                             '6. Оноо < 85% бол хариуцагчид засах ажил үүсгэж мэдэгдэнэ'],
           ['Дараах нөхцөл', 'Аудитын бүртгэл, бүсийн оноо, шаардлагатай бол засах ажил үүснэ'],
           ['Алтернатив урсгал', '5а. Сүлжээгүй бол утсанд хадгалж, сүлжээ сэргэхэд илгээнэ\n'
                                 '6а. Тухайн аудитаас ажил үүссэн бол давхар үүсгэхгүй']]
table(s, M, Cm(3.7), [Cm(4.4), Cm(15.4)], uc_rows, size=13, head=False, bold_cols=(0,),
      heights=[Cm(h) for h in (1.0, 1.0, 1.0, 1.0, 3.7, 1.0, 1.35)],
      fills={0: {'body': CARD}, 1: {'body': WHITE}}, zebra=False)
box = rect(s, M, Cm(15.0), Cm(19.8), Cm(2.5), CARD, radius=Cm(0.2))
shape_text(box, [[('Бусад 8 use-case ', {'bold': True}),
                  ('ижил загвараар — тоглогч, өмнөх ба дараах нөхцөл, үндсэн ба алтернатив урсгалтай — дипломын '
                   '2-р бүлэгт тодорхойлогдсон', {})]], size=12.5, color=INK, align=PP_ALIGN.LEFT)
x = Cm(22.0)
w = W - M - x
label(s, x, Cm(3.75), w, 'UC-02-ын мөрдөлт')
links = [('Шаардлага', 'FR-07 QR-ээр аудит · FR-08 аудитаас засах ажил · FR-09 улаан шошго · FR-18 сүлжээгүй горим'),
         ('Алгоритм', 'Оноо тооцох, засах ажил үүсгэх (слайд 17) · outbox дараалал (слайд 18)'),
         ('Диаграмм', 'Дарааллын диаграмм (нэмэлт слайд Н2) · аудитын урсгал (слайд 17)'),
         ('Тест', 'Оноо < 85% бол засах ажил нэг удаа үүснэ; дахин илгээхэд давхардахгүй')]
for i, (head, body) in enumerate(links):
    y = Cm(4.45 + i * 3.25)
    rect(s, x, y, w, Cm(2.95), CARD, radius=Cm(0.2))
    label(s, x + Cm(0.5), y + Cm(0.35), w - Cm(1.0), head, color=NAVY)
    text(s, x + Cm(0.5), y + Cm(0.95), w - Cm(1.0), Cm(1.9), body, size=12.5)

# --------------------------------------------------------------------------- 13. Мөрдөх матриц
s = slide('Шаардлагын мөрдөх матриц', 3,
          'Мөрдөх матриц шаардлага бүрийг use-case ба хэрэгжилтийн гурван түвшинтэй холбоно: кодонд хэрэгжсэн, '
          'тестээр шалгасан, туршилтаар хэмжсэн. 22 шаардлагын 20 нь кодонд бүрэн хэрэгжиж, автомат тестээр '
          'шалгагдсан. Сүлжээгүй горим бичих үйлдэлд л, KPI самбар хэсэгчлэн хэрэгжсэн; AI тоймыг хиймэл хариугаар '
          'тестэлсэн. Туршилтын баганыг III үзлэгийн туршилтын ажиллагаагаар бөглөнө.')
trace = [('FR-01', 'Нэвтрэх, эрхийн хүрээ', 'UC-01', 'y', 'y'), ('FR-02', 'Урилга, эрх олгох', '–', 'y', 'y'),
         ('FR-03', 'Ажил үүсгэх, оноох', 'UC-03', 'y', 'y'), ('FR-04', 'Ач холбогдлоор харах', 'UC-03', 'y', 'y'),
         ('FR-05', 'Оноолт, өглөөний мэдэгдэл', 'UC-03', 'y', 'y'),
         ('FR-06', 'Талбай, бүс, стандарт', 'UC-09', 'y', 'y'), ('FR-07', 'QR-ээр аудит', 'UC-02', 'y', 'y'),
         ('FR-08', 'Аудитаас засах ажил', 'UC-02', 'y', 'y'), ('FR-09', 'Улаан шошго, зураг', 'UC-02', 'y', 'y'),
         ('FR-10', 'Өдрийн бүртгэл', 'UC-07', 'y', 'y'), ('FR-11', 'Долоо хоногийн тайлан', 'UC-07', 'y', 'y'),
         ('FR-12', 'Сарын тайлан, хаалт', 'UC-04', 'y', 'y'), ('FR-13', 'Хагас жил, жилийн тайлан', 'UC-04', 'y', 'y'),
         ('FR-14', 'AI тойм бичвэр', 'UC-04', 'y', 'p'), ('FR-15', 'Өглөөний хурлын самбар', 'UC-08', 'y', 'y'),
         ('FR-16', 'Gemba явалт', 'UC-05', 'y', 'y'), ('FR-17', 'Сайжруулалтын санаа', 'UC-06', 'y', 'y'),
         ('FR-18', 'Сүлжээгүй горим', 'UC-02,05,07', 'p', 'y'), ('FR-19', 'Монгол, англи хэл', 'бүгд', 'y', 'y'),
         ('FR-20', 'Үйл ажиллагааны түүх', 'UC-04', 'y', 'y'), ('FR-21', 'Бүртгэл устгах хүсэлт', '–', 'y', 'y'),
         ('FR-22', 'KPI самбар (K1–K6)', 'UC-04,08', 'p', 'n')]
half_w = (CW - Cm(0.6)) / 2
cols = [Cm(1.7), Cm(5.9), Cm(2.75), Cm(1.5), Cm(1.5)]
cols.append(int(half_w) - sum(cols))
for k in range(2):
    rows = [['FR', 'Шаардлага', 'Use-case', 'Код', 'Тест', 'Турш.']]
    rows += [[a, b, c, d, e, 'n'] for a, b, c, d, e in trace[k * 11:(k + 1) * 11]]
    table(s, M + int(k * (half_w + Cm(0.6))), Cm(3.6), cols, rows, size=11.5, row_h=Cm(0.94), bold_cols=(0,),
          align={2: PP_ALIGN.CENTER, 3: PP_ALIGN.CENTER, 4: PP_ALIGN.CENTER, 5: PP_ALIGN.CENTER})
stats = [('Кодонд хэрэгжсэн', '20 бүрэн · 2 хэсэгчлэн', 'FR-18 (офлайн унших), FR-22 (KPI самбар)'),
         ('Тестээр шалгасан', '20 бүрэн · 1 хэсэгчлэн · 1 үгүй', 'FR-14 хиймэл AI хариугаар, FR-22 дараа'),
         ('Туршилтаар хэмжсэн', 'III үзлэгт', 'Туршилтын ажиллагаа: 2 + 4 долоо хоног')]
stat_w = (CW - Cm(1.0)) / 3
for i, (head, value, note) in enumerate(stats):
    x = M + int(i * (stat_w + Cm(0.5)))
    rect(s, x, Cm(15.3), int(stat_w), Cm(2.2), CARD, radius=Cm(0.2))
    label(s, x + Cm(0.5), Cm(15.5), int(stat_w) - Cm(1.0), head, size=9.5)
    text(s, x + Cm(0.5), Cm(16.0), int(stat_w) - Cm(1.0), Cm(0.7), value, size=14, bold=True, color=NAVY)
    text(s, x + Cm(0.5), Cm(16.72), int(stat_w) - Cm(1.0), Cm(0.7), note, size=10.5, color=MUTED)

# --------------------------------------------------------------------------- 14. Архитектур
s = slide('Системийн архитектур', 4,
          'Давхаргат архитектур: вэб ба гар утасны хоёр клиент нэг REST API ашиглана, бизнесийн дүрэм нэг газар. '
          'API-д JWT нэвтрэлт ба эрхийн шалгалт, controller, service, repository давхаргууд, мөн сануулга, аудитын '
          'хуваарь, сар хаах cron ажлууд бий. Өгөгдөл PostgreSQL-д, 25 хүснэгт, 38 migration-оор удирдагдана. '
          'AI үйлчилгээ сонголтоор, зөвхөн нэргүй тоо илгээнэ.')
picture(s, 'architecture.png', M, Cm(3.5), Cm(19.9), Cm(14.1))
x = Cm(22.0)
w = W - M - x
layers = [('Клиент', 'React 18 · Vite · Tailwind (вэб)\nFlutter (Android) · сүлжээгүй дараалал'),
          ('API', 'NestJS 11 · REST · JWT + дүрийн эрх\nCron: сануулга, аудит, сар хаах'),
          ('Өгөгдөл', 'PostgreSQL · TypeORM\n25 хүснэгт · 38 migration'),
          ('Чанар', '830 + 1 017 + 116 автомат тест\nWCAG 2.1 AA · монгол / англи')]
for i, (head, body) in enumerate(layers):
    y = Cm(3.75 + i * 3.45)
    rect(s, x, y, w, Cm(3.15), NAVY if i == 1 else CARD, radius=Cm(0.25))
    label(s, x + Cm(0.6), y + Cm(0.45), w - Cm(1.2), head, color=YELLOW if i == 1 else MUTED)
    text(s, x + Cm(0.6), y + Cm(1.1), w - Cm(1.2), Cm(1.9), body.split('\n'), size=13,
         color=WHITE if i == 1 else INK, after=2)

# --------------------------------------------------------------------------- 15. Байршуулалт
s = slide('Байршуулалтын диаграм', 4,
          'Системийг Docker Compose-оор аль ч Linux сервер дээр байршуулна. nginx статик вэбийг үйлчилж, /api '
          'хүсэлтийг backend рүү дамжуулна. Гар утас шууд API-тай HTTPS-ээр холбогдоно. Redis кэш, нөөцлөлт, '
          'Prometheus, Grafana мониторинг сонголтоор асна. Нууц тохиргоо .env файлд, git-д орохгүй.')
picture(s, 'deployment.png', M, Cm(3.5), Cm(21.4), Cm(14.1))
x = Cm(23.6)
w = W - M - x
label(s, x, Cm(3.9), w, 'Docker Compose')
text(s, x, Cm(4.6), w, Cm(12.8),
     [[('Үндсэн: ', {'bold': True}), ('postgres 15, backend (NestJS :3000), admin-web (nginx: статик файл, '
                                    '/api → backend)', {})],
      [('Сонголтоор: ', {'bold': True}), ('redis (кэш), backup, prometheus + grafana', {})],
      [('Өгөгдөл: ', {'bold': True}), ('postgres_data_prod, attachments_prod volume', {})],
      [('Нууцлал: ', {'bold': True}), ('тохиргоо .env файлд, git-д орохгүй; HTTPS', {})],
      [('NFR-08: ', {'bold': True}), ('аль ч Linux сервер дээр нэг командаар асна', {})]],
     size=13.5, after=12, line=1.1)

# --------------------------------------------------------------------------- 16. Өгөгдлийн сан ба класс
s = slide('Өгөгдлийн сан ба класс диаграм', 4,
          'Зүүн талд өгөгдлийн сангийн гол хүснэгтүүд, баруун талд шинжилгээний класс диаграмм. Нийт 25 хүснэгт, '
          'бүх хүснэгт байгууллагаар тусгаарлагдана. Ажил нь аудит, Gemba явалт, санаанаас үүсэх бөгөөд эх '
          'сурвалжаа хадгална; үүгээр нэг эх сурвалжаас ажил давхар үүсэхээс сэргийлнэ.')
for k, (name, caption) in enumerate([('er.png', 'ER диаграм: гол хүснэгтүүд (нийт 25 хүснэгт, 38 migration)'),
                                     ('class_analysis.png', 'Шинжилгээний класс диаграм: ажил аудит, Gemba, '
                                                            'санаанаас үүснэ')]):
    x = M + int(k * (half_w + Cm(0.6)))
    rect(s, x, Cm(3.6), int(half_w), Cm(12.9), WHITE, radius=Cm(0.25), line=LINE)
    picture(s, name, x + Cm(0.3), Cm(3.85), int(half_w) - Cm(0.6), Cm(12.4))
    text(s, x, Cm(16.75), int(half_w), Cm(0.7), caption, size=11.5, color=MUTED, italic=True,
         align=PP_ALIGN.CENTER)

# --------------------------------------------------------------------------- 17. Алгоритм: аудит
s = slide('Алгоритмын зохиомж: аудитын оноо ба засах ажил', 4,
          'Гол алгоритм: аудитын оноо = авсан онооны нийлбэр / боломжит онооны нийлбэр × 100. Хариулаагүй асуулт 0, '
          'текст хариулт нотолгоо тул оноонд орохгүй. 85%-иас доош бол бүсийн хариуцагчид 7 хоногийн хугацаатай '
          'засах ажил үүснэ, 70%-иас доош бол өндөр ач холбогдолтой. Нэг аудитаас нэг л ажил үүснэ. Хугацааны '
          'төвөгшил асуултын тоогоор O(n).')
picture(s, 'audit_flow.png', M, Cm(3.4), Cm(11.4), Cm(14.4))
x = Cm(13.7)
w = W - M - x
label(s, x, Cm(3.75), w, 'Оноо тооцох')
text(s, x, Cm(4.4), w, Cm(1.2), 'S = round( Σ авсан / Σ боломжит × 100 )', size=22, bold=True, color=NAVY)
text(s, x, Cm(5.75), w, Cm(1.3), 'Асуулт бүрийн авсан оноо ба боломжит дээд оноо; хариулаагүй асуулт 0, текст '
     'хариулт нотолгоо тул оноонд орохгүй', size=12, color=MUTED)
rules = [('S ≥ 85%', 'тэнцсэн', GREEN), ('S < 85%', 'засах ажил, 7 хоног', AMBER), ('S < 70%', 'өндөр ач холбогдол',
                                                                                   RED)]
rule_w = (w - Cm(0.8)) / 3
for i, (cond, result, colour) in enumerate(rules):
    bx = x + int(i * (rule_w + Cm(0.4)))
    box = rect(s, bx, Cm(7.2), int(rule_w), Cm(1.6), WHITE, radius=Cm(0.2), line=colour, line_w=Pt(1.5))
    shape_text(box, [[(cond, {'bold': True, 'color': colour})], [(result, {})]], size=12.5, color=INK)
code(s, x, Cm(9.25), w, Cm(5.75), '''
function correctiveWork(run):
    if run.status == "draft" or run.score >= 85:
        return null                          // тэнцсэн эсвэл ноорог
    zone = zoneOf(run)
    return createTask(
        title    = "5S follow-up: " + zone.name,
        assignee = zone.owner,
        due      = today + 7 days,
        priority = run.score < 70 ? high : medium,
        source   = (audit_run, run.id))      // нэг аудитаас нэг ажил
''', size=11.5)
text(s, x, Cm(15.4), w, Cm(1.5),
     [[('Төвөгшил: ', {'bold': True}), ('O(n), n — асуултын тоо.  ', {}), ('Давхардалгүй: ', {'bold': True}),
       ('(sourceType, sourceId) хосоор шалгана, дахин илгээсэн аудит шинэ ажил үүсгэхгүй.', {})]],
     size=12, color=INK)

# --------------------------------------------------------------------------- 18. Алгоритм: outbox, KPI, сануулга
s = slide('Алгоритмын зохиомж: сүлжээгүй горим, KPI, сануулга', 4,
          'Сүлжээгүй горимд серверт огт хүрээгүй хүсэлтийг дараалалд хадгалж, сүлжээ сэргэхэд хийсэн дарааллаар нь '
          'илгээнэ; давхар бүртгэлээс сэргийлж timeout-ыг дахин илгээхгүй. KPI-г ажил, аудитын бүртгэлээс нэг '
          'дамжилтаар O(n+m) хугацаанд тооцно. Өдрийн сануулга байгууллагын цагийн бүсээр, нэг хүнд өдөрт нэг л удаа '
          'очно. Сарын тайлан хаагдахад хуулбар хадгалж, нээлттэй ба хаагдсан тайлан нэг дүрмээр тоологдоно.')
label(s, M, Cm(3.75), half_w, 'Сүлжээгүй горимын дараалал (outbox)')
picture(s, 'outbox.png', M, Cm(4.4), int(half_w), Cm(6.6), border=LINE, valign='top')
text(s, M, Cm(11.55), int(half_w), Cm(6.0),
     ['•  Зөвхөн серверт огт хүрээгүй (холболтгүй) хүсэлтийг дараалалд хийнэ; timeout-ыг дахин илгээхгүй',
      '•  Сүлжээ сэргэх, апп нээгдэх, «Одоо илгээх» үед хийсэн дарааллаар нь илгээнэ',
      '•  4xx → бичилтийг хасч хэрэглэгчид мэдэгдэнэ; 5xx, 401 → хадгалаад дараа дахин'],
     size=13, after=7)
x = M + int(half_w + Cm(0.6))
label(s, x, Cm(3.75), half_w, 'Бүтээмжийн үзүүлэлт (K1–K4)')
code(s, x, Cm(4.4), int(half_w), Cm(5.3), '''
function kpis(tasks, audits, period, tz, today):   // O(n + m)
    due    = tasks: status != backlog, dueDate in period
    onTime = due: localDate(completedAt, tz) <= dueDate
    K1 = 100 * |onTime| / |due|
    late = tasks: not done, dueDate < today
    K2 = (|late|, mean(today - dueDate))
    K3 = (mean(audit.score), share of zones >= 85)
    K4 = mean(completedAt - createdAt)  // аудитаас үүссэн ажил
    return K1, K2, K3, K4
''', size=10.5)
items = [('Өдрийн сануулга', 'Cron цаг тутам; байгууллагын цагийн бүсээр тооцно; өдрийг эх сурвалж болгосон тул нэг '
                             'хүнд өдөрт нэг л сануулга (idempotent)'),
         ('Сарын тайлан', 'Хаагдахад хуулбар хадгална; биелэлтийг ажлын одоогийн төлвөөс бус, хэзээ дууссанаас '
                          'тооцно'),
         ('Ажлыг бүлэглэх', 'хоцорсон → өнөөдөр → энэ долоо хоног → дараа → хугацаагүй → дууссан, O(n)')]
for i, (head, body) in enumerate(items):
    y = Cm(10.3 + i * 2.35)
    text(s, x, y, int(half_w), Cm(2.1), [[(head + ': ', {'bold': True, 'color': NAVY}), (body, {})]], size=13,
         line=1.05)

# --------------------------------------------------------------------------- 19. Интерфейс
s = slide('Хэрэглэгчийн интерфейс', 4,
          'Вэб дээрх өглөөний хурлын самбар: өчигдөр, өнөөдөр, анхаарах ажил, 5S, багаас гарсан санаа, Gemba-гийн '
          'биелэлтийг нэг дэлгэцэд харуулна. Гар утсанд 5S талбай, бүсийн оноо, шалгах хуудас, улаан шошго; сүлжээгүй '
          'үед ч бүртгэл хадгалагдана. Интерфейс монгол, англи хэлтэй, цайвар ба бараан горимтой.')
picture(s, 'ui_huddle.png', M, Cm(3.65), Cm(19.8), Cm(12.5), border=LINE, valign='top')
text(s, M, Cm(16.35), Cm(19.8), Cm(1.2),
     'Вэб: өглөөний хурлын самбар — өчигдөр · өнөөдөр · анхаарах · 5S · санаа · Gemba', size=11.5, color=MUTED,
     italic=True)
x = Cm(22.0)
w1 = phone(s, 'mob_fives.png', x, Cm(3.65), Cm(10.0))
phone(s, 'mob_audit.png', x + w1 + Cm(0.45), Cm(3.65), Cm(10.0))
text(s, x, Cm(14.1), W - M - x, Cm(1.2), 'Гар утас: 5S талбай → бүсийн аудит, улаан шошго (сүлжээгүй үед ч)',
     size=11.5, color=MUTED, italic=True)
text(s, x, Cm(15.4), W - M - x, Cm(2.2),
     [[('Монгол / англи', {'bold': True}), (' · цайвар, бараан горим', {})],
      [('WCAG 2.1 AA', {'bold': True}), (' · том үсэг, нэг гараар ашиглах', {})]], size=12.5, after=4)

# --------------------------------------------------------------------------- 20. Төлөв, төлөвлөгөө
s = slide('Одоогийн төлөв ба цаашдын төлөвлөгөө', 4,
          'Ажлын төлвийг гурван түвшинд ялгаж харууллаа: кодонд хэрэгжсэн, тестээр шалгасан, туршилтаар хэмжсэн. '
          'Гол урсгалууд кодонд хэрэгжиж, автомат тестээр баталгаажсан; KPI самбар хэсэгчлэн, хэрэглэгчийн судалгаа '
          'хийгдээгүй. II үзлэгт KPI самбар, серверт байршуулалт, судалгаа; III үзлэгт туршилтын ажиллагаа, SUS '
          'үнэлгээ, KPI-г өмнө–дараа загвараар хэмжинэ.')
status = [['Хэсэг', 'Код', 'Тест', 'Туршилт'],
          ['Ажил, төсөл, мэдэгдэл', 'y', 'y', 'n'], ['5S талбай, аудит, засах ажил', 'y', 'y', 'n'],
          ['Өдрийн ба долоо хоногийн тайлан', 'y', 'y', 'n'], ['Сарын тайлан, хаалт, архив', 'y', 'y', 'n'],
          ['Gemba, санаа, өглөөний хурал', 'y', 'y', 'n'], ['Сүлжээгүй горим (бичих)', 'p', 'y', 'n'],
          ['KPI самбар (K1–K6)', 'p', 'n', 'n'], ['Хэрэглэгчийн судалгаа, SUS', 'n', 'n', 'n']]
table(s, M, Cm(3.65), [Cm(9.4), Cm(2.8), Cm(2.8), Cm(2.8)], status, size=12.5, row_h=Cm(0.95),
      align={1: PP_ALIGN.CENTER, 2: PP_ALIGN.CENTER, 3: PP_ALIGN.CENTER})
tiles = [('830', 'backend тест'), ('1 017', 'вэб тест'), ('116', 'гар утасны тест'), ('39', 'e2e (18 + 21 бодит)')]
tile_w = (Cm(17.8) - Cm(0.9)) / 4
for i, (num, name) in enumerate(tiles):
    x = M + int(i * (tile_w + Cm(0.3)))
    rect(s, x, Cm(12.85), int(tile_w), Cm(2.45), CARD, radius=Cm(0.2))
    text(s, x, Cm(13.1), int(tile_w), Cm(1.1), num, size=24, bold=True, color=NAVY, align=PP_ALIGN.CENTER)
    text(s, x, Cm(14.25), int(tile_w), Cm(0.8), name, size=10.5, color=MUTED, align=PP_ALIGN.CENTER)
text(s, M, Cm(15.7), Cm(17.8), Cm(1.8),
     [[('✓', {'font': SYMBOL, 'color': GREEN, 'bold': True}), (' бүрэн   ', {}),
       ('∼', {'font': SYMBOL, 'color': AMBER, 'bold': True}), (' хэсэгчлэн   ', {}),
       ('–', {'color': GREY, 'bold': True}), (' хараахан үгүй.  Автомат тест бүгд амжилттай (2026.09.29).', {})]],
     size=11, color=MUTED)
x = Cm(20.6)
w = W - M - x
label(s, x, Cm(3.75), w, 'Үзлэгийн хуваарь')
rect(s, x + Cm(0.3), Cm(4.85), Cm(0.08), Cm(9.3), LINE)
milestones = [('I үзлэг · 09.30', 'Онол, аргазүй, шаардлага, use-case, архитектур, алгоритм, ном зүй', GREEN),
              ('II үзлэг · 11.03–05', 'Алгоритмын кодчилол, тест: KPI самбар (FR-22), серверт байршуулах, '
                                     'хэрэглэгчийн судалгаа', YELLOW),
              ('III үзлэг · 12.02–03', 'Туршилтын ажиллагаа (2 + 4 долоо хоног), K1–K7, SUS, дүгнэлт', WHITE),
              ('Урьдчилсан хамгаалалт · 12.28', 'Илтгэл, эцсийн засвар', WHITE)]
for i, (head, body, colour) in enumerate(milestones):
    y = Cm(4.5 + i * 3.1)
    rect(s, x + Cm(0.04), y + Cm(0.05), Cm(0.6), Cm(0.6), colour, line=NAVY if colour == WHITE else None,
         line_w=Pt(1.5), shape=MSO_SHAPE.OVAL)
    text(s, x + Cm(1.1), y, w - Cm(1.1), Cm(0.8), head, size=15, bold=True, color=INK)
    text(s, x + Cm(1.1), y + Cm(0.85), w - Cm(1.1), Cm(2.0), body, size=12, color=MUTED)

# --------------------------------------------------------------------------- 21. Ном зүй
s = slide('Ном зүй', 4,
          'Ном зүйд нийт 31 эх сурвалж: lean, 5S, бүтээмжийн хэмжилтийн сонгодог бүтээлүүд; программ хангамжийн '
          'инженерчлэл, архитектурын ном; ISO/IEC 25010, WCAG, JWT стандарт; SUS үнэлгээний судалгаа; ижил төстэй '
          'системүүд ба технологийн баримт бичиг.')
# Дугаар нь дипломын ном зүйтэй ижил (biblatex sorting=none: эхлээд иш татсан дарааллаар)
refs_left = [(1, 'Hirano, H. (1995). 5 Pillars of the Visual Workplace. Productivity Press.'),
             (2, 'Imai, M. (1997). Gemba Kaizen: A Commonsense, Low-Cost Approach to Management. McGraw-Hill.'),
             (3, 'Liker, J. K. (2004). The Toyota Way. McGraw-Hill.'),
             (4, 'Ohno, T. (1988). Toyota Production System. Productivity Press.'),
             (5, 'Womack, J. P., Jones, D. T. (1996). Lean Thinking. Simon & Schuster.'),
             (6, 'Mann, D. (2005). Creating a Lean Culture. Productivity Press.'),
             (7, 'Osada, T. (1991). The 5S\'s: Five Keys to a Total Quality Environment. APO.'),
             (8, 'Doerr, J. (2018). Measure What Matters. Portfolio/Penguin.'),
             (9, 'Deming, W. E. (1986). Out of the Crisis. MIT Press.'),
             (10, 'OECD (2001). Measuring Productivity: OECD Manual. OECD Publishing.'),
             (11, 'Parmenter, D. (2015). Key Performance Indicators (3rd ed.). Wiley.')]
refs_right = [(12, 'Brooke, J. (1996). SUS: A "quick and dirty" usability scale. Taylor & Francis.'),
              (13, 'Bangor, A., Kortum, P. T., Miller, J. T. (2008). An Empirical Evaluation of the SUS. IJHCI, 24(6).'),
              (19, 'Sommerville, I. (2016). Software Engineering (10th ed.). Pearson.'),
              (20, 'Beck, K. (2002). Test-Driven Development: By Example. Addison-Wesley.'),
              (21, 'ISO/IEC 25010:2011. System and software quality models.'),
              (25, 'Jones, M., Bradley, J., Sakimura, N. (2015). JSON Web Token (JWT), RFC 7519. IETF.'),
              (27, 'W3C (2018). Web Content Accessibility Guidelines (WCAG) 2.1.'),
              (28, 'Cockburn, A. (2001). Writing Effective Use Cases. Addison-Wesley.'),
              (29, 'Fowler, M. (2002). Patterns of Enterprise Application Architecture. Addison-Wesley.'),
              (30, 'Martin, R. C. (2017). Clean Architecture. Prentice Hall.')]
for k, (head, refs) in enumerate([('Бүтээмж, lean, хэмжилт', refs_left),
                                  ('Программ хангамж, стандарт, үнэлгээ', refs_right)]):
    x = M + int(k * (half_w + Cm(0.6)))
    label(s, x, Cm(3.7), half_w, head)
    text(s, x, Cm(4.4), int(half_w), Cm(11.4),
         [[('[%d]  ' % n, {'bold': True, 'color': NAVY}), (r, {})] for n, r in refs], size=11.5, after=4.5, line=1.0)
rect(s, M, Cm(14.9), CW, Cm(2.6), CARD, radius=Cm(0.2))
text(s, M + Cm(0.5), Cm(15.2), CW - Cm(1.0), Cm(2.1),
     [[('[14–18]  ', {'bold': True, 'color': NAVY}), ('SafetyCulture, Weekdone, Asana, KaiNexus, Tervene', {})],
      [('[22–24, 26, 31]  ', {'bold': True, 'color': NAVY}),
       ('React, NestJS, PostgreSQL, Flutter, Claude API-ийн баримт бичиг (2026.09.29-нд хандсан)', {})],
      [('Нийт 31 эх сурвалж; дугаар нь дипломын Ном зүйтэй ижил.', {'color': MUTED})]], size=11.5, after=3)

# --------------------------------------------------------------------------- 22. Баярлалаа
s = dark_slide('Анхаарал хандуулсанд баярлалаа. Асуулт, санал хүлээн авахад бэлэн байна. Нэмэлт слайдад эдийн '
               'засгийн урьдчилсан тооцоо, дарааллын болон үйл ажиллагааны диаграмм бий.')
rect(s, 0, 0, W, Cm(0.35), YELLOW)
label(s, M, Cm(5.2), CW, 'Дипломын ажлын I үзлэг', color=YELLOW, size=12, align=PP_ALIGN.CENTER)
text(s, M, Cm(6.3), CW, Cm(3.6), ['Анхаарал хандуулсанд', 'баярлалаа'], size=44, bold=True, color=WHITE,
     align=PP_ALIGN.CENTER, after=0, line=1.0)
text(s, M, Cm(10.6), CW, Cm(1.0), 'Асуулт, санал?', size=20, color=PALE, align=PP_ALIGN.CENTER)
text(s, M, Cm(15.6), CW, Cm(1.6), [AUTHOR + ' · B190910014', 'Удирдагч: Доктор (Ph.D), дэд профессор Г.Ганчимэг'],
     size=12, color=PALE, align=PP_ALIGN.CENTER, after=3)

# --------------------------------------------------------------------------- Нэмэлт слайдууд
s = slide('Эдийн засгийн урьдчилсан тооцоо', 0,
          'Эдийн засгийн тооцоо нь таамаглалд суурилсан: 25 ажилтантай (5 менежер) жишээ байгууллага, цагийн өртөг '
          '12 000 төгрөг, жилд 3 408 цаг хэмнэнэ. NPV эерэг, нөхөх хугацаа нэг жил хүрэхгүй. Туршилтын үед K5 болон '
          'бусад үзүүлэлтээр бодит хэмнэлтийг хэмжиж шинэчилнэ.', number='Н1')
econ = [['Үзүүлэлт', 'Үр дүн'], ['Анхны хөрөнгө оруулалт (Inv)', '24 000 000 ₮'],
        ['Жилийн өгөөж (3 408 цаг хэмнэлт)', '40 896 000 ₮'], ['Жилийн зардал', '6 900 000 ₮'],
        ['Жилийн цэвэр мөнгөн урсгал', '33 996 000 ₮'], ['NPV (5 жил, r = 12%)', '98 547 972 ₮'],
        ['Нөхөх хугацаа', '≈ 0.79 жил (9.5 сар)'], ['Ашигт ажиллагаа (Pa)', '83.13%'],
        ['Нийт бүтээмж (TP)', '5.93']]
table(s, M, Cm(3.7), [Cm(10.6), Cm(6.4)], econ, size=13, row_h=Cm(1.05), align={1: PP_ALIGN.RIGHT})
text(s, M, Cm(13.6), Cm(17.0), Cm(2.6),
     ['NPV = Σ [ FV / (1 + r)^i ] − Inv,   i = 1…5',
      'Нөхөх хугацаа = Inv / PV(1) = 24.0 / 30.35 ≈ 0.79 жил'], size=16, color=NAVY, after=8)
x = Cm(20.0)
w = W - M - x
rect(s, x, Cm(3.7), w, Cm(9.5), CARD, radius=Cm(0.25))
label(s, x + Cm(0.6), Cm(4.3), w - Cm(1.2), 'Таамаглал')
text(s, x + Cm(0.6), Cm(5.0), w - Cm(1.2), Cm(8.0),
     ['•  25 ажилтан (5 менежер), 12 000 ₮/цаг', '•  Ажилтан 2 цаг/долоо хоног хэмнэнэ',
      '•  Менежер 3 цаг/долоо хоног хэмнэнэ', '•  Сарын тайлан 24 цаг → автомат', '•  n = 5 жил, r = 12%'],
     size=13.5, after=8)
box = rect(s, x, Cm(13.6), w, Cm(3.9), NAVY, radius=Cm(0.25))
shape_text(box, 'Тооцоо таамаглалд суурилсан; туршилтын ажиллагаанд K5 ба бусад үзүүлэлтээр бодит хэмнэлтийг '
                'хэмжиж шинэчилнэ.', size=13, color=WHITE, align=PP_ALIGN.LEFT)

s = slide('Дарааллын диаграм: 5S аудит (UC-02)', 0,
          'Гар утаснаас аудит илгээхэд API оноог тооцож хадгална; оноо 85%-иас доош бол засах ажил үүсч, бүсийн '
          'эзэнд мэдэгдэл очно. Сүлжээгүй бол бичилт утсанд хүлээгдэнэ.', number='Н2')
picture(s, 'seq_audit.png', M, Cm(3.5), CW, Cm(14.1))
s = slide('Дарааллын диаграм: сарын тайлан хаах (UC-04)', 0,
          'Сар хаах өдөр болоход цагийн хуваарь тухайн сарын бүртгэлийг хуулбарлан хадгалж, менежер тойм бичвэрийг '
          'бичиж эсвэл AI-аар ноороглуулаад батална.', number='Н3')
picture(s, 'seq_close.png', M, Cm(3.5), CW, Cm(14.1))
s = slide('Үйл ажиллагааны диаграм: ажлын амьдралын мөчлөг', 0,
          'Ажил оноогдохоос дуусах хүртэл менежер, систем, ажилтан гэсэн гурван оролцогчийн үйл ажиллагаа.',
          number='Н4')
picture(s, 'activity_task.png', M, Cm(3.5), CW, Cm(14.1))
s = slide('Класс диаграм (зохиомжийн түвшин)', 0,
          'Зохиомжийн түвшний класс диаграмм: controller, service, repository давхаргууд ба тэдгээрийн хамаарал.',
          number='Н5')
picture(s, 'class_design.png', M, Cm(3.5), CW, Cm(14.1))

out = os.path.join(HERE, 'output', 'Uzleg1_iltgel.pptx')
prs.save(out)
print(out, len(prs.slides._sldIdLst), 'slides')
