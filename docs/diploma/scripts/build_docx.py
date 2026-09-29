# -*- coding: utf-8 -*-
"""Дипломын бичвэрийг Word (.docx) болгох. python build_docx.py"""
import os

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

import content as C

HERE = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
FIG = os.path.join(HERE, 'fig')
FONT = 'Times New Roman'

doc = Document()

# --- Хуудасны тохиргоо: А4, зүүн 3 см, баруун 1.5 см, дээд доод 2 см
for section in doc.sections:
    section.page_width, section.page_height = Cm(21), Cm(29.7)
    section.left_margin, section.right_margin = Cm(3), Cm(1.5)
    section.top_margin, section.bottom_margin = Cm(2), Cm(2)


def set_font(style, size, bold=False):
    style.font.name = FONT
    style.font.size = Pt(size)
    style.font.bold = bold
    style.font.color.rgb = RGBColor(0, 0, 0)
    rpr = style.element.get_or_add_rPr()
    fonts = rpr.find(qn('w:rFonts'))
    if fonts is None:
        fonts = OxmlElement('w:rFonts')
        rpr.append(fonts)
    for attr in ('w:ascii', 'w:hAnsi', 'w:cs', 'w:eastAsia'):
        fonts.set(qn(attr), FONT)
    # Гарчгийн загвар theme фонт (Calibri Light) заадаг; түүнийг арилгана
    for attr in ('w:asciiTheme', 'w:hAnsiTheme', 'w:cstheme', 'w:eastAsiaTheme'):
        if fonts.get(qn(attr)) is not None:
            del fonts.attrib[qn(attr)]


normal = doc.styles['Normal']
set_font(normal, 12)
normal.paragraph_format.line_spacing = 1.5
normal.paragraph_format.space_after = Pt(0)

for name, size in (('Heading 1', 14), ('Heading 2', 13), ('Heading 3', 12)):
    style = doc.styles[name]
    set_font(style, size, bold=True)
    style.font.italic = False
    style.paragraph_format.space_before = Pt(12 if name != 'Heading 1' else 0)
    style.paragraph_format.space_after = Pt(6)
    style.paragraph_format.line_spacing = 1.15
    style.paragraph_format.keep_with_next = True
doc.styles['Heading 1'].paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER


def para(text='', align=None, bold=False, size=None, indent=True, space_after=0, italic=False):
    p = doc.add_paragraph()
    if text:
        run = p.add_run(text)
        run.bold = bold
        run.italic = italic
        if size:
            run.font.size = Pt(size)
    p.paragraph_format.alignment = align if align is not None else WD_ALIGN_PARAGRAPH.JUSTIFY
    if indent and align is None:
        p.paragraph_format.first_line_indent = Cm(1.25)
    p.paragraph_format.space_after = Pt(space_after)
    return p


def field(paragraph, instruction):
    run = paragraph.add_run()
    for tag, text in (('begin', None), (None, instruction), ('separate', None), (None, ''), ('end', None)):
        if tag:
            el = OxmlElement('w:fldChar')
            el.set(qn('w:fldCharType'), tag)
        else:
            el = OxmlElement('w:instrText' if text == instruction else 'w:t')
            el.set(qn('xml:space'), 'preserve')
            el.text = text if text != '' else 'Агуулгыг шинэчлэхийн тулд энд хулганы баруун товч → Update Field дарна уу.'
        run._r.append(el)


def page_break():
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)


def shade(cell, color):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), color)
    tc_pr.append(shd)


def cell_text(cell, text, bold=False, size=10, align=WD_ALIGN_PARAGRAPH.LEFT):
    cell.text = ''
    lines = text.split('\n')
    p = cell.paragraphs[0]
    for i, line in enumerate(lines):
        if i:
            p = cell.add_paragraph()
        run = p.add_run(line)
        run.bold = bold
        run.font.size = Pt(size)
        p.paragraph_format.line_spacing = 1.0
        p.paragraph_format.alignment = align
        p.paragraph_format.space_after = Pt(1)


def caption(text, before_table=True):
    p = para(text, align=WD_ALIGN_PARAGRAPH.RIGHT if before_table else WD_ALIGN_PARAGRAPH.CENTER,
             bold=before_table, size=11, indent=False, space_after=3)
    p.paragraph_format.keep_with_next = before_table
    p.paragraph_format.space_before = Pt(6)
    return p


def table(title, header, rows, widths=None):
    caption(title)
    t = doc.add_table(rows=1, cols=len(header))
    t.style = 'Table Grid'
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, h in enumerate(header):
        cell_text(t.rows[0].cells[i], h, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
        shade(t.rows[0].cells[i], 'D9E2F3')
    for row in rows:
        cells = t.add_row().cells
        for i, value in enumerate(row):
            centred = value in ('✓', '—', 'хэсэгчлэн', 'M', 'S', 'C')
            cell_text(cells[i], value, align=WD_ALIGN_PARAGRAPH.CENTER if centred else WD_ALIGN_PARAGRAPH.LEFT)
    if widths:
        for row in t.rows:
            for i, w in enumerate(widths):
                row.cells[i].width = Cm(w)
    # Толгой мөр хуудас бүрт давтагдана
    tr_pr = t.rows[0]._tr.get_or_add_trPr()
    el = OxmlElement('w:tblHeader')
    el.set(qn('w:val'), 'true')
    tr_pr.append(el)
    para('', indent=False)


def usecase_table(title, rows):
    caption(title)
    t = doc.add_table(rows=0, cols=2)
    t.style = 'Table Grid'
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for key, value in rows:
        cells = t.add_row().cells
        cell_text(cells[0], key, bold=True)
        shade(cells[0], 'EEF2F7')
        cell_text(cells[1], value)
        cells[0].width, cells[1].width = Cm(3.6), Cm(12.4)
    para('', indent=False)


def figure(name, text, width):
    p = doc.add_paragraph()
    p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.keep_with_next = True
    p.add_run().add_picture(os.path.join(FIG, name), width=Cm(width))
    caption(text, before_table=False)


def code(text):
    t = doc.add_table(rows=1, cols=1)
    t.style = 'Table Grid'
    cell = t.rows[0].cells[0]
    shade(cell, 'F5F5F5')
    cell.text = ''
    for i, line in enumerate(text.split('\n')):
        p = cell.paragraphs[0] if i == 0 else cell.add_paragraph()
        run = p.add_run(line)
        run.font.name = 'Consolas'
        run._element.rPr.rFonts.set(qn('w:eastAsia'), 'Consolas')
        run.font.size = Pt(9)
        p.paragraph_format.line_spacing = 1.0
        p.paragraph_format.space_after = Pt(0)
    para('', indent=False)


# ------------------------------------------------------------------ Нүүр хуудас
for line in ('ШИНЖЛЭХ УХААН ТЕХНОЛОГИЙН ИХ СУРГУУЛЬ', 'МЭДЭЭЛЭЛ, ХОЛБООНЫ ТЕХНОЛОГИЙН СУРГУУЛЬ',
             'КОМПЬЮТЕРИЙН УХААНЫ ТЭНХИМ'):
    para(line, align=WD_ALIGN_PARAGRAPH.CENTER, bold=True, size=13, indent=False)
for _ in range(5):
    para('', indent=False)
para('[Оюутны овог, нэр]', align=WD_ALIGN_PARAGRAPH.CENTER, bold=True, size=14, indent=False)
para('', indent=False)
para(C.TITLE.upper(), align=WD_ALIGN_PARAGRAPH.CENTER, bold=True, size=16, indent=False)
para(C.TITLE_EN, align=WD_ALIGN_PARAGRAPH.CENTER, size=12, indent=False, italic=True)
para('', indent=False)
para('Бакалаврын дипломын төсөл (F.CS370)', align=WD_ALIGN_PARAGRAPH.CENTER, size=13, indent=False)
para('Программ хангамжийн инженерчлэл', align=WD_ALIGN_PARAGRAPH.CENTER, size=13, indent=False)
for _ in range(5):
    para('', indent=False)
para('Удирдагч багш: [цол, овог нэр] ..................', align=WD_ALIGN_PARAGRAPH.RIGHT, size=12, indent=False)
para('Зөвлөгч багш: [цол, овог нэр] ..................', align=WD_ALIGN_PARAGRAPH.RIGHT,
     size=12, indent=False)
for _ in range(4):
    para('', indent=False)
para('Улаанбаатар хот, 2026 он', align=WD_ALIGN_PARAGRAPH.CENTER, bold=True, size=12, indent=False)

# ------------------------------------------------------------------ Хураангуй, агуулга
page_break()
doc.add_heading('ХУРААНГУЙ', level=1)
for text in C.ABSTRACT:
    para(text)
page_break()
doc.add_heading('АГУУЛГА', level=1)
field(doc.add_paragraph(), 'TOC \\o "1-3" \\h \\z \\u')
page_break()
doc.add_heading('ТОВЧИЛСОН ҮГИЙН ЖАГСААЛТ', level=1)
t = doc.add_table(rows=0, cols=2)
for short, long in C.ABBREVIATIONS:
    cells = t.add_row().cells
    cell_text(cells[0], short, bold=True, size=12)
    cell_text(cells[1], long, size=12)
    cells[0].width, cells[1].width = Cm(2.5), Cm(13.5)
page_break()

# ------------------------------------------------------------------ Их бие
for item in C.BODY:
    kind = item[0]
    if kind == 'h1':
        doc.add_heading(item[1], level=1)
    elif kind == 'h2':
        doc.add_heading(item[1], level=2)
    elif kind == 'h3':
        doc.add_heading(item[1], level=3)
    elif kind == 'p':
        para(item[1])
    elif kind == 'bullets':
        for text in item[1]:
            p = doc.add_paragraph(style='List Bullet')
            p.add_run(text)
            p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    elif kind == 'table':
        table(item[1], item[2], item[3], item[4] if len(item) > 4 else None)
    elif kind == 'uc':
        usecase_table(item[1], item[2])
    elif kind == 'fig':
        figure(item[1], item[2], item[3])
    elif kind == 'code':
        code(item[1])
    elif kind == 'pagebreak':
        page_break()

# ------------------------------------------------------------------ Ном зүй
doc.add_heading('НОМ ЗҮЙ', level=1)
for i, ref in enumerate(C.REFERENCES, 1):
    p = para(f'[{i}] {ref}', indent=False)
    p.paragraph_format.left_indent = Cm(1)
    p.paragraph_format.first_line_indent = Cm(-1)
    p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT

# ------------------------------------------------------------------ Хуудасны дугаар
footer = doc.sections[0].footer.paragraphs[0]
footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
field(footer, 'PAGE')
for el in footer._p.iter(qn('w:t')):
    el.text = ''
doc.sections[0].different_first_page_header_footer = True

# Нээхэд агуулга, хуудасны дугаарыг Word өөрөө шинэчилнэ
settings = doc.settings.element
update = OxmlElement('w:updateFields')
update.set(qn('w:val'), 'true')
settings.append(update)

out = os.path.join(HERE, 'output', 'archive', 'Diplom_Uzleg1.docx')
doc.save(out)
print(out)
