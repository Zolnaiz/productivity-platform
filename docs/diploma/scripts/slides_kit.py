# -*- coding: utf-8 -*-
"""Үзлэгийн илтгэлүүдийн (build_slides.py, build_slides2.py) нийтлэг хэсэг: өнгө, үсэг, хуудасны
хэмжээ, бичвэр, хүснэгт, зураг, код зурах туслах функцууд, мөн гарчигтай ба бараан слайд.

Нэг загвар дээр хоёр илтгэл бүтээдэг тул өнгө, хэмжээг зөвхөн энд өөрчилнө.
"""
import os

from lxml import etree
from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Cm, Emu, Pt

HERE = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
FIG = os.path.join(HERE, 'latex', 'Figures')
OUTPUT = os.path.join(HERE, 'output')

NAVY = RGBColor(0x16, 0x20, 0x2C)
NAVY2 = RGBColor(0x1E, 0x2B, 0x39)
YELLOW = RGBColor(0xF5, 0xC5, 0x18)
INK = RGBColor(0x20, 0x2A, 0x35)
MUTED = RGBColor(0x65, 0x71, 0x7D)
CARD = RGBColor(0xF3, 0xF5, 0xF7)
LINE = RGBColor(0xD8, 0xDE, 0xE4)
GREEN = RGBColor(0x2F, 0x7D, 0x68)
AMBER = RGBColor(0xA8, 0x6B, 0x0C)
RED = RGBColor(0xB8, 0x50, 0x42)
BLUE = RGBColor(0x2F, 0x5D, 0xA8)
GREY = RGBColor(0x9A, 0xA5, 0xB1)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
PALE = RGBColor(0xCB, 0xD3, 0xDA)
TINT = RGBColor(0xFE, 0xF6, 0xD8)
CODE_INK = RGBColor(0xE6, 0xED, 0xF3)
CODE_NOTE = RGBColor(0x8B, 0x98, 0xA5)
FONT = 'Arial'
MONO = 'Consolas'
SYMBOL = 'Segoe UI Symbol'

AUTHOR = 'Б.Билгүүн'
TITLE = 'Байгууллагын бүтээмжийг дэмжих платформ хөгжүүлэлт'

W, H = Cm(33.867), Cm(19.05)
M = Cm(1.6)
CW = W - 2 * M


def new_deck():
    prs = Presentation()
    prs.slide_width, prs.slide_height = W, H
    prs.core_properties.title = TITLE
    prs.core_properties.author = AUTHOR
    return prs


# --------------------------------------------------------------------------- helpers

def text(s, x, y, w, h, lines, size=16, color=INK, bold=False, align=PP_ALIGN.LEFT, font=FONT,
         spacing=None, after=4, line=1.08, anchor=MSO_ANCHOR.TOP, italic=False):
    """Бичвэр. lines: мөрүүд; мөр бүр str, (str, style) эсвэл тэдгээрийн жагсаалт (run-ууд)."""
    box = s.shapes.add_textbox(int(x), int(y), int(w), int(h))
    frame = box.text_frame
    frame.word_wrap = True
    frame.margin_left = frame.margin_right = frame.margin_top = frame.margin_bottom = 0
    frame.vertical_anchor = anchor
    if not isinstance(lines, list):
        lines = [lines]
    for i, item in enumerate(lines):
        p = frame.paragraphs[0] if i == 0 else frame.add_paragraph()
        p.alignment = align
        p.space_after = Pt(after)
        p.line_spacing = line
        for part in (item if isinstance(item, list) else [item]):
            value, style = part if isinstance(part, tuple) else (part, {})
            run = p.add_run()
            run.text = value
            f = run.font
            f.name = style.get('font', font)
            f.size = Pt(style.get('size', size))
            f.bold = style.get('bold', bold)
            f.italic = style.get('italic', italic)
            f.color.rgb = style.get('color', color)
            gap = style.get('spacing', spacing)
            if gap:
                run._r.get_or_add_rPr().set('spc', str(int(gap * 100)))
    return box


def rect(s, x, y, w, h, fill, radius=None, line=None, line_w=Pt(1), shape=None):
    kind = shape or (MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE)
    x, y, w, h = int(x), int(y), int(w), int(h)
    box = s.shapes.add_shape(kind, x, y, w, h)
    if radius and kind == MSO_SHAPE.ROUNDED_RECTANGLE:
        box.adjustments[0] = min(0.5, radius / min(w, h))
    if fill is None:
        box.fill.background()
    else:
        box.fill.solid()
        box.fill.fore_color.rgb = fill
    if line is None:
        box.line.fill.background()
    else:
        box.line.color.rgb = line
        box.line.width = line_w
    box.shadow.inherit = False
    return box


def label(s, x, y, w, value, color=MUTED, size=10.5, align=PP_ALIGN.LEFT):
    """Жижиг том үсгэн гарчиг (eyebrow)."""
    return text(s, x, y, w, Cm(0.6), value.upper(), size=size, bold=True, color=color, spacing=1.2, align=align)


def shape_text(box, value, size=14, color=INK, bold=False, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE):
    frame = box.text_frame
    frame.word_wrap = True
    frame.margin_left = frame.margin_right = Cm(0.25)
    frame.margin_top = frame.margin_bottom = Cm(0.1)
    frame.vertical_anchor = anchor
    lines = value if isinstance(value, list) else [value]
    for i, item in enumerate(lines):
        p = frame.paragraphs[0] if i == 0 else frame.add_paragraph()
        p.alignment = align
        for part in (item if isinstance(item, list) else [item]):
            v, style = part if isinstance(part, tuple) else (part, {})
            run = p.add_run()
            run.text = v
            run.font.name = style.get('font', FONT)
            run.font.size = Pt(style.get('size', size))
            run.font.bold = style.get('bold', bold)
            run.font.color.rgb = style.get('color', color)
            if style.get('spacing'):
                run._r.get_or_add_rPr().set('spc', str(int(style['spacing'] * 100)))


def picture(s, name, x, y, w, h, border=None, valign='middle'):
    """Зургийг хайрцагт багтааж, харьцааг нь хадгалан голлуулна."""
    path = os.path.join(FIG, name)
    iw, ih = Image.open(path).size
    x, y, w, h = int(x), int(y), int(w), int(h)
    if iw / ih > w / h:
        pw, ph = w, int(w * ih / iw)
    else:
        pw, ph = int(h * iw / ih), h
    top = y if valign == 'top' else y + (h - ph) // 2
    pic = s.shapes.add_picture(path, x + (w - pw) // 2, top, pw, ph)
    if border is not None:
        pic.line.color.rgb = border
        pic.line.width = Pt(0.75)
    return pic


def phone(s, name, x, y, h):
    """Гар утасны дэлгэцийн зургийг хүрээтэй харуулна; өргөнийг буцаана."""
    path = os.path.join(FIG, name)
    iw, ih = Image.open(path).size
    pad = Cm(0.2)
    x, y, h = int(x), int(y), int(h)
    inner_h = h - 2 * pad
    inner_w = int(inner_h * iw / ih)
    rect(s, x, y, inner_w + 2 * pad, h, NAVY, radius=Cm(0.45))
    s.shapes.add_picture(path, x + pad, y + pad, inner_w, inner_h)
    return inner_w + 2 * pad


def code(s, x, y, w, h, source, size=11):
    """Псевдо код: бараан дэвсгэр, // тайлбарыг бүдэг өнгөөр."""
    rect(s, x, y, w, h, NAVY, radius=Cm(0.25))
    lines = []
    for raw in source.strip('\n').split('\n'):
        body, _, note = raw.partition('//')
        runs = [(body, {'color': CODE_INK})]
        if note:
            runs.append(('//' + note, {'color': CODE_NOTE}))
        lines.append(runs)
    return text(s, x + Cm(0.45), y + Cm(0.35), w - Cm(0.9), h - Cm(0.7), lines, size=size, font=MONO,
                after=0, line=1.0)


MARKS = {'y': ('✓', GREEN), 'p': ('∼', AMBER), 'n': ('–', GREY)}


def _border(cell, sides=('B',), color='D8DEE4', width=9525):
    tc_pr = cell._tc.get_or_add_tcPr()
    for side in reversed(('L', 'R', 'T', 'B')):
        if side not in sides:
            continue
        ln = etree.Element(qn('a:ln' + side), w=str(width), cap='flat', cmpd='sng', algn='ctr')
        fill = etree.SubElement(ln, qn('a:solidFill'))
        etree.SubElement(fill, qn('a:srgbClr'), val=color)
        etree.SubElement(ln, qn('a:prstDash'), val='solid')
        tc_pr.insert(0, ln)


def table(s, x, y, widths, rows, size=12, row_h=Cm(0.9), head=True, align=None, fills=None, bold_cols=(),
          head_fill=NAVY, head_color=WHITE, zebra=True, anchor=MSO_ANCHOR.MIDDLE, heights=None):
    """Хүснэгт: толгой мөр бараан, мөрүүд ээлжлэн цайвар, доод зураас нимгэн.

    Олон мөртэй нүдэнд heights-ээр өндрийг нь заана: PowerPoint мөрийг өөрөө сунгадаг ч Canva
    импортлохдоо хадгалсан өндрийг ашиглаж, илүү мөрийг тасалдаг.
    """
    heights = heights or [row_h] * len(rows)
    frame = s.shapes.add_table(len(rows), len(widths), x, y, sum(widths), sum(heights))
    tbl = frame.table
    pr = tbl._tbl.tblPr
    pr.set('firstRow', '0')
    pr.set('bandRow', '0')
    style = pr.find(qn('a:tableStyleId'))
    if style is None:
        style = etree.SubElement(pr, qn('a:tableStyleId'))
    style.text = '{2D5ABB26-0587-4C30-8999-92F81FD0307C}'  # No Style, No Grid
    for j, w in enumerate(widths):
        tbl.columns[j].width = w
    for i, row in enumerate(rows):
        tbl.rows[i].height = heights[i]
        is_head = head and i == 0
        for j, value in enumerate(row):
            cell = tbl.cell(i, j)
            cell.margin_left = cell.margin_right = Cm(0.18)
            cell.margin_top = cell.margin_bottom = Cm(0.06)
            cell.vertical_anchor = anchor
            fill = (fills or {}).get(j)
            if isinstance(fill, dict):
                fill = fill.get('head' if is_head else 'body')
            if fill is None:
                fill = head_fill if is_head else (CARD if zebra and i % 2 == 0 else WHITE)
            cell.fill.solid()
            cell.fill.fore_color.rgb = fill
            if not is_head:
                _border(cell)
            frame_text = cell.text_frame
            frame_text.word_wrap = True
            mark = MARKS.get(value) if isinstance(value, str) and not is_head else None
            paragraphs = [value] if mark else str(value).split('\n')
            for k, line_value in enumerate(paragraphs):
                p = frame_text.paragraphs[0] if k == 0 else frame_text.add_paragraph()
                p.alignment = (align or {}).get(j, PP_ALIGN.LEFT) if not mark else PP_ALIGN.CENTER
                p.space_after = Pt(1)
                run = p.add_run()
                if mark:
                    run.text, colour = mark
                    run.font.name = SYMBOL
                    run.font.size = Pt(size + 2)
                    run.font.bold = True
                    run.font.color.rgb = colour
                    continue
                run.text = line_value
                run.font.name = FONT
                run.font.size = Pt(size)
                run.font.bold = is_head or j in bold_cols
                run.font.color.rgb = head_color if is_head else INK
    return tbl


# --------------------------------------------------------------------------- slides

def add_slide(prs, title, section, footer, notes, number=None):
    """Гарчигтай слайд: дээр хэсгийн нэр, шар зураас, гарчиг, дугаар; доор зохиогч ба үзлэг."""
    s = prs.slides.add_slide(prs.slide_layouts[6])
    label(s, M, Cm(0.85), Cm(24), section)
    rect(s, M, Cm(1.5), Cm(1.1), Cm(0.13), YELLOW)
    text(s, M, Cm(1.72), Cm(28), Cm(1.4), title, size=28, bold=True, color=NAVY)
    n = number or '%02d' % len(prs.slides._sldIdLst)
    text(s, W - M - Cm(3), Cm(0.8), Cm(3), Cm(0.7), n, size=13, bold=True, color=NAVY, align=PP_ALIGN.RIGHT)
    rect(s, M, Cm(18.05), CW, Emu(9525), LINE)
    text(s, M, Cm(18.2), Cm(22), Cm(0.5), AUTHOR + ' · ' + TITLE, size=9, color=MUTED)
    text(s, W - M - Cm(10), Cm(18.2), Cm(10), Cm(0.5), footer, size=9, color=MUTED, align=PP_ALIGN.RIGHT)
    s.notes_slide.notes_text_frame.text = notes
    return s


def add_dark_slide(prs, notes):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    rect(s, 0, 0, W, H, NAVY)
    s.notes_slide.notes_text_frame.text = notes
    return s
