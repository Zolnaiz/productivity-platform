# -*- coding: utf-8 -*-
"""Сургуулийн "Төсөл, дипломын ажлын даалгавар" маягтыг бөглөнө.
python scripts/fill_task_form.py -> output/Diplom_daalgavar.docx"""
import os

from docx import Document

HERE = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
SRC = os.path.join(HERE, 'materials', "Temlate for project's requirement-ПХ-2026-төсөл.docx")
OUT = os.path.join(HERE, 'output', 'Diplom_daalgavar.docx')

doc = Document(SRC)


def set_text(paragraph, text):
    """Эхний run-ий хэлбэрийг хадгалж, бичвэрийг солино."""
    runs = paragraph.runs
    if not runs:
        paragraph.add_run(text)
        return
    runs[0].text = text
    for run in runs[1:]:
        run.text = ''


def fill(prefix, value):
    for p in doc.paragraphs:
        if p.text.strip().startswith(prefix):
            label = p.text.split(':')[0].strip()
            set_text(p, f'{label}: {value}')
            return
    raise SystemExit(f'not found: {prefix}')


for p in doc.paragraphs:
    if p.text.strip().startswith('Оюутны овог'):
        set_text(p, 'Оюутны овог нэр: Бямбадоржийн Билгүүн\tКод: B190910014')
        break

fill('Мэргэжлийн код, нэр', '071405000000002306, ПРОГРАММ ХАНГАМЖИЙН ИНЖЕНЕРЧЛЭЛ')
fill('Хичээлийн код, нэр', 'Бакалаврын дипломын төсөл (F.CS370)')
fill('Улирал', '2026-2027 оны хичээлийн жилийн намар')
fill('Сэдвийн нэр монголоор', 'Байгууллагын бүтээмжийг дэмжих платформ хөгжүүлэлт')
fill('Сэдвийн нэр англиар', 'Development of a Platform to Support Organizational Productivity')
fill('Удирдагчийн овог, нэр', 'Доктор (Ph.D), дэд профессор Г.Ганчимэг, F.CS05')
fill('Ажлын зорилго',
     'Байгууллагын 5S, бүтээмжийн хэрэгслүүд болон ажлын удирдлагыг нэгтгэж, менежерт ажлыг '
     'бодит хугацаанд оноож хянах, ажилтанд ажлаа гар утаснаас хөтлөх, тайланг автоматаар '
     'гаргах боломж олгох вэб болон гар утасны платформыг зохион бүтээж, хэрэгжүүлэх.')

items = [
    'Онол, аргазүйн судалгаа: 5S, Gemba, Kaizen, PDCA, шатлалт өдөр тутмын хурал; бүтээмжийг хэмжих '
    'үзүүлэлт (KPI) ба «өмнө–дараа» үнэлгээний загвар; ижил төстэй системийн харьцуулалт (SafetyCulture, '
    'Weekdone, Asana, KaiNexus, Tervene), SWOT шинжилгээ.',
    'Шинжилгээ, зохиомж: функциональ ба функциональ бус шаардлага, use-case диаграм ба тодорхойлолт, '
    'шаардлагын мөрдөх матриц, класс, дараалал, үйл ажиллагааны диаграм, давхаргат архитектур, '
    'өгөгдлийн сангийн схем; аудитын оноо, сүлжээгүй горим, сануулга, сарын тайлан, бүтээмжийн '
    'үзүүлэлтийн алгоритмын зохиомж.',
    'Хэрэгжүүлэлт, туршилт: NestJS + PostgreSQL API, React вэб апп, Flutter гар утасны апп; '
    'автомат тест, WCAG 2.1 AA хүртээмж; туршилтын байгууллагад ажиллуулж KPI-г хэмжих, SUS аргаар '
    'үнэлэх; эдийн засгийн үр ашгийн тооцоо.',
]
dots = [p for p in doc.paragraphs if p.style.name == 'List Paragraph' and set(p.text.strip()) <= {'.'} and p.text.strip()]
for p, text in zip(dots, items):
    set_text(p, text)

schedule = {
    'Үзлэг I': ('2026-09-30', 'Онол, аргазүйн судалгаа, шаардлага, use-case, архитектур, алгоритмын зохиомж, ном зүй'),
    'Үзлэг II': ('2026-11-03..05', 'Алгоритмын кодчилол, шинжилгээ, туршилт, үр дүн'),
    'Үзлэг III': ('2026-12-02, 03', 'Системийн хөгжүүлэлт, туршилт, үр дүн, дүгнэлт, хавсралт'),
    'Хамгаалалт': ('2026-12-28 (урьдчилсан)', 'Илтгэл, диплом (хэвлэмэл), туршилтын программ'),
}
table = doc.tables[0]
for row in table.rows[1:]:
    key = row.cells[0].text.strip()
    if key in schedule:
        date, result = schedule[key]
        row.cells[1].text = date
        row.cells[2].text = result

os.makedirs(os.path.dirname(OUT), exist_ok=True)
doc.save(OUT)
print(OUT)
